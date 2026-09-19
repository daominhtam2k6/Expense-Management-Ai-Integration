"""Update the previously delivered Word template with observed live results."""
from copy import deepcopy
from pathlib import Path
import shutil
import zipfile
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
source = Path(r'C:\Users\LENOVO\Downloads\05_GenAI_SoftwareDevelopment_functional-testing_completed.docx')
backup = source.with_name(source.stem+'_before_live_20260919.docx')
if not backup.exists():
    shutil.copy2(source, backup)
doc = Document(source)

def replace(p, value):
    if p.runs:
        p.runs[0].text=value
        for r in p.runs[1:]:
            r.text=''
    else:
        p.add_run(value)

def cell(c, value):
    replace(c.paragraphs[0], value)
    for p in c.paragraphs[1:]:
        replace(p, '')

replace(doc.paragraphs[15], '3. Báo cáo kết quả test (Test report) – cập nhật 19/09/2026')
replace(doc.paragraphs[16],
    'Kết luận: CHƯA ĐẠT TOÀN BỘ. Backend ngày 19/09: 71 passed, 4 failed, 22 subtests passed, '
    '7 cảnh báo; hai lỗi xác nhận ở AC-018 và AC-019. Frontend/build ngày 18/09: 50/50 đạt, '
    'coverage statements 62,86%, branches 46,55%, functions 57,68%, lines 67,90%; PWA 35 entries. '
    'Ngày 19/09: Gemini live có phản hồi; Resend từ chối khóa; web production không kết nối được; '
    '9 kiểm tra trình duyệt local đạt. Các kết quả Pass trong suite cũ chỉ xác nhận nhánh đã chạy, '
    'không xác nhận mọi điều kiện trong ca tổng hợp. Chưa có nghiệm thu thủ công của thành viên nhóm. '
    'PostgreSQL, Tauri, tải vận hành, TLS production, nhận email thực tế và retention backup chưa được kiểm chứng.')

report=doc.tables[3]
for row in report.rows[1:]:
    test_id=row.cells[0].text
    cell(row.cells[1], '18–19/09/2026')
    cell(row.cells[2], 'Tự động; chưa duyệt thủ công')
    if test_id in ['FT-06','FT-08']:
        cell(row.cells[3], 'Fail')
        cell(row.cells[4], 'Vừa')
        cell(row.cells[5], 'Tên Food/food/ Food được tạo trùng cùng user/loại.' if test_id=='FT-06' else 'Schema nhận năm 1999 và 2101, ngoài khoảng 2000–2100.')
        cell(row.cells[6], 'Tái hiện 19/09; tests/test_baseline_requirements.py; chưa sửa code.')
    else:
        cell(row.cells[3], 'Đạt phần đã chạy')
        cell(row.cells[5], 'Suite cũ đạt; chưa chứng minh đầy đủ mọi nhánh trong ca tổng hợp.')

entries=[
 ('LIVE-01','Gemini thật','Gọi adapter ứng dụng đến Gemini, không mock.','Khóa và model từ môi trường.','Số liệu giả: thu 1.000.000, chi 200.000 VND.','Có phản hồi văn bản; không gửi dữ liệu cá nhân.','PASS','12,22 giây; 196 ký tự. Không chứng nhận độ đúng mọi câu trả lời hoặc SLA.','Không có'),
 ('LIVE-02','Resend thật','Gọi GET /domains bằng khóa cấu hình.','Khóa Resend trong môi trường.','Không gửi thư hay nội dung tài khoản.','API chấp nhận khóa hợp lệ.','FAIL','HTTP 400, validation_error: API key is invalid. Chưa gửi/nhận email.','Cao'),
 ('LIVE-03','Web production','GET /, /login, /api/health, /api/ready, /api/unknown-live-test.','Tên miền từ FRONTEND_URL.','https://daominhtam2k6.dpdns.org','HTTP 200 cho route hợp lệ, JSON 404 cho API không tồn tại; HTTPS hợp lệ.','BLOCKED','Cả 5 request ConnectionError; DNS không trả IP trong kiểm tra local. Chưa kết luận TLS/health.','Chưa xác định'),
 ('LIVE-04','Trình duyệt đăng ký/đăng nhập','Dùng Edge headless thao tác form thật đến API local thật.','Server riêng, SQLite tạm, không mock API.','Tài khoản ngẫu nhiên @example.com; không gửi email.','Đăng ký về login; đăng nhập về dashboard.','PASS','2 luồng đạt; dữ liệu thử được dọn cùng database tạm.','Không có'),
 ('LIVE-05','7 trang nghiệp vụ','Mở dashboard, transactions, categories, budgets, goals, reports, assistant.','Đã đăng nhập trên instance thử.','Dữ liệu tài khoản mới; service worker bị chặn trong browser test.','Trang hiện, không có API HTTP >=400 hoặc pageerror.','PASS','7/7 smoke đạt; chưa thực hiện tất cả CRUD qua UI hoặc cache/offline PWA.','Không có'),
 ('LIVE-06','PostgreSQL','Kiểm tra điều kiện chạy database production-like.','Cần Docker daemon/PostgreSQL thử.','docker info; cấu hình hiện dùng SQLite.','Có môi trường PostgreSQL để kiểm tra migration.','BLOCKED','Docker Desktop Linux Engine không kết nối được. SQLite migration sạch đã đạt 18/09, không thay thế PostgreSQL.','Chưa xác định'),
 ('LIVE-07','Tauri/tải/backup','Kiểm tra phạm vi còn thiếu.','Cần toolchain và môi trường vận hành phù hợp.','Cargo chưa tìm thấy; chưa có dữ liệu vòng đời backup.','Bộ cài chạy; tải đạt AC-015; backup đạt AC-021.','NOT RUN','Chưa kiểm thử thực tế; không đánh dấu đạt.','Chưa xác định'),
]
for tid,feature,desc,pre,data,expected,status,observed,severity in entries:
    for table,values in [(doc.tables[2],[tid,feature,desc,pre,data,expected,'Live check 19/09/2026']),
                         (report,[tid,'19/09/2026','Tự động; chưa duyệt thủ công',status,severity,observed,'Bằng chứng: deliverables/live-*-results-20260919.json; docs/test-report.md'])]:
        table._tbl.append(deepcopy(table.rows[1]._tr))
        for c,v in zip(table.rows[-1].cells,values):
            cell(c,v)
doc.save(source)
output=ROOT/'deliverables/05_functional_testing_live_20260919.docx'
doc.save(output)
for path in [source,output]:
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
    check=Document(path)
    assert len(check.tables)==4 and len(check.tables[2].rows)==26 and len(check.tables[3].rows)==26
    print(path)
print('Backup:',backup)
