from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r"C:\Users\LENOVO\Downloads\Tai lieu du an\04_GenAI_SoftwareDevelopment_object-oriented-design_class-specification.docx")
OUTPUT = ROOT / "docs/outputs/04_GenAI_SoftwareDevelopment_object-oriented-design_class-specification_with_method_diagrams.docx"
DIAGRAM_DIR = ROOT / "docs/diagrams/uml/class-method-activities"


def set_font(run):
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


def slug_for(operation: str) -> str:
    import re
    return re.sub(r"[^a-z0-9]+", "-", operation.lower()).strip("-")


document = Document(SOURCE)

# Explain the as-built ownership boundary without changing the document's class grouping.
for paragraph in document.paragraphs:
    if paragraph.text.strip().startswith("Phần này đặc tả 9 lớp thực thể"):
        paragraph.clear()
        paragraph.add_run(
            "Phần này đặc tả 9 entity ORM và các thao tác Router/Core liên quan. "
            "Các thao tác được nhóm theo entity nghiệp vụ để truy vết, nhưng không được coi là "
            "phương thức thành viên đã cài đặt của entity nếu mã nguồn đặt chúng tại Router/Core."
        )
    elif paragraph.text.strip() == "b) Các phương thức":
        paragraph.clear()
        paragraph.add_run("b) Các thao tác Router/Core liên quan và biểu đồ hoạt động")
    elif paragraph.text.strip().startswith("Phương thức "):
        text = paragraph.text.strip().replace("Phương thức ", "Thao tác ", 1)
        paragraph.clear()
        paragraph.add_run(text)

inserted = 0
for table in list(document.tables):
    if len(table.columns) != 2 or len(table.rows) < 7 or table.cell(0, 0).text.strip() != "Tên":
        continue
    operation = table.cell(0, 1).text.strip()
    image_path = DIAGRAM_DIR / f"{slug_for(operation)}.png"
    if not image_path.exists():
        raise FileNotFoundError(image_path)

    table.cell(0, 0).text = "Thao tác (Router/Core)"
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = True
    caption.add_run(f"Biểu đồ hoạt động – {operation}")
    picture = document.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    with Image.open(image_path) as image:
        pixel_width, pixel_height = image.size
    width_cm = min(13.5, 14.0 * pixel_width / pixel_height)
    picture.add_run().add_picture(str(image_path), width=Cm(width_cm))
    table._tbl.addnext(caption._p)
    caption._p.addnext(picture._p)
    inserted += 1

for style in document.styles:
    if style.type == 1:
        style.font.name = "Times New Roman"
        style.font.size = Pt(13)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
for paragraph in document.paragraphs:
    for run in paragraph.runs:
        set_font(run)
for table in document.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_font(run)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
document.save(OUTPUT)
print(f"{OUTPUT}\nInserted diagrams: {inserted}")
