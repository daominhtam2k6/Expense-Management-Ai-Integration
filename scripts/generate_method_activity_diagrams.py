import re
import subprocess
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r"C:\Users\LENOVO\Downloads\Tai lieu du an\04_GenAI_SoftwareDevelopment_object-oriented-design_class-specification.docx")
OUT_DIR = ROOT / "docs/diagrams/uml/class-method-activities"


def clean(text: str) -> str:
    return " ".join(text.replace("\n", " ").split())


def plantuml_text(text: str, width: int = 54) -> str:
    text = (
        clean(text)
        .replace("\\", "\\\\")
        .replace(":", " -")
        .replace("→", "–")
    )
    words = text.split()
    lines, current = [], []
    for word in words:
        if current and len(" ".join(current + [word])) > width:
            lines.append(" ".join(current))
            current = [word]
        else:
            current.append(word)
    if current:
        lines.append(" ".join(current))
    return "\\n".join(lines)


def slugify(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


document = Document(SOURCE)
OUT_DIR.mkdir(parents=True, exist_ok=True)
records = []

for table in document.tables:
    if len(table.columns) != 2 or len(table.rows) < 7:
        continue
    if table.cell(0, 0).text.strip() != "Tên":
        continue
    values = {row.cells[0].text.strip(): row.cells[1].text.strip() for row in table.rows}
    operation = clean(values.get("Tên", ""))
    if not operation:
        continue
    steps = [clean(item) for item in re.split(r"\s*\d+\)\s*", values.get("Luồng xử lý", "")) if clean(item)]
    slug = slugify(operation)
    lines = [
        "@startuml",
        f"title BIỂU ĐỒ HOẠT ĐỘNG — {operation}",
        "skinparam backgroundColor white",
        "skinparam defaultFontName Arial",
        "skinparam defaultFontSize 15",
        "skinparam dpi 300",
        "skinparam shadowing false",
        "skinparam activity {",
        "  BackgroundColor #F4FBF8",
        "  BorderColor #299D78",
        "  ArrowColor #345B52",
        "}",
        "start",
    ]
    start_condition = clean(values.get("Điều kiện bắt đầu", ""))
    if start_condition:
        lines.append(f":Điều kiện bắt đầu\\n{plantuml_text(start_condition)}; ")
    for step in steps:
        lines.append(f":{plantuml_text(step)};")
    end_condition = clean(values.get("Điều kiện kết thúc", ""))
    if end_condition:
        lines.append(f":Kết quả\\n{plantuml_text(end_condition)};")
    lines.extend(["stop", "@enduml", ""])
    (OUT_DIR / f"{slug}.puml").write_text("\n".join(lines), encoding="utf-8")
    records.append((operation, slug))

(OUT_DIR / "manifest.tsv").write_text(
    "operation\tslug\n" + "\n".join(f"{operation}\t{slug}" for operation, slug in records) + "\n",
    encoding="utf-8",
)
(OUT_DIR / "render-notes.md").write_text(
    "# Activity Diagram cho thao tác Router/Core\n\n"
    f"- Nguồn: `{SOURCE}` và đặc tả thao tác trong 40 bảng phương thức.\n"
    "- Trạng thái: mô tả luồng as-built đã được ghi trong tài liệu; không coi hàm Router/Core là phương thức thành viên của entity ORM.\n"
    "- Kiểu sơ đồ: UML Activity Diagram, một hình cho mỗi thao tác.\n"
    "- Font ảnh: Arial 15, nền trắng, PNG/SVG render bằng PlantUML local.\n"
    "- Nội dung mỗi hình: điều kiện bắt đầu, các bước xử lý và kết quả/điều kiện kết thúc.\n",
    encoding="utf-8",
)
print(f"Generated {len(records)} PlantUML sources in {OUT_DIR}")
