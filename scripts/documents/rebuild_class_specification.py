"""Replace Part 2 of the supplied OOD document with complete class specifications."""
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parents[2]

SOURCE = Path(r"C:\Users\LENOVO\Downloads\Tai lieu du an\04_GenAI_SoftwareDevelopment_object-oriented-design_presentation_chuanhoa.docx")
OUTPUT = (PROJECT_ROOT / "docs/archive/ban-khac/04_GenAI_SoftwareDevelopment_object-oriented-design_class-specification.docx")
BLUE, LIGHT_BLUE, LIGHT_GREEN = "2E74B5", "D9E5F6", "E6F4EE"


def io(name, type_, size): return (name, type_, size)
def method(name, desc, inputs, outputs, flow, pre, post):
    return {"name": name, "desc": desc, "inputs": inputs, "outputs": outputs, "flow": flow, "pre": pre, "post": post}


SPECS = [
    {
        "name": "User", "table": "users", "desc": "Quản lý tài khoản, hồ sơ, xác thực, avatar và khôi phục mật khẩu.",
        "attrs": [
            ("id", "String/UUID", "36 ký tự", "Khóa chính; sinh tự động."),
            ("username", "String", "3–50 ký tự", "Tên đăng nhập; duy nhất, có index."),
            ("email", "String/Email", "Tối đa 254 ký tự", "Email đăng nhập và khôi phục; duy nhất, có index."),
            ("display_name", "String?", "Tối đa 80 ký tự", "Tên hiển thị; có thể rỗng."),
            ("password_hash", "String", "Theo chuỗi bcrypt", "Mật khẩu đã băm; không lưu mật khẩu rõ."),
            ("avatar_url", "String?", "Độ dài URL", "Đường dẫn ảnh đại diện; có thể rỗng."),
            ("created_at", "DateTime(TZ)", "Ngày giờ + múi giờ", "Thời điểm tạo tài khoản."),
            ("reset_token_hash", "String?", "64 ký tự SHA-256", "Hash token đặt lại mật khẩu."),
            ("reset_token_expiry", "DateTime(TZ)?", "Ngày giờ + múi giờ", "Thời điểm token hết hiệu lực."),
        ],
        "methods": [
            method("register", "Tạo tài khoản người dùng mới.", [io("username","String","3–50 ký tự"),io("email","Email","≤254 ký tự"),io("password","String","≥6 ký tự")], [io("user","UserOut","5 trường công khai")], "1) Chuẩn hóa username/email. 2) Kiểm tra trùng. 3) Băm mật khẩu. 4) Lưu User và trả dữ liệu công khai.", "Chưa đăng nhập; dữ liệu đầu vào hợp lệ.", "Tài khoản được commit; password_hash tồn tại; mật khẩu rõ không được lưu."),
            method("login", "Xác thực bằng username hoặc email và cấp JWT.", [io("identifier","String","Không rỗng"),io("password","String","≥6 ký tự")], [io("access_token","JWT String","Theo token ký"),io("token_type","String","6 ký tự: bearer")], "1) Tìm User theo username, sau đó email. 2) So khớp bcrypt. 3) Tạo JWT chứa user id. 4) Trả token; sai thì trả 401.", "Tài khoản tồn tại; client chưa có phiên hợp lệ.", "Đúng: token hợp lệ được trả. Sai: không tạo phiên và trả thông báo lỗi."),
            method("get_profile", "Lấy thông tin tài khoản hiện tại.", [io("access_token","JWT String","Theo token ký")], [io("user","UserOut","5 trường công khai")], "1) Giải mã JWT. 2) Tìm User theo sub. 3) Trả hồ sơ không chứa password_hash/token reset.", "JWT còn hiệu lực.", "Trả đúng người dùng hiện tại hoặc 401."),
            method("update_profile", "Cập nhật tên hiển thị, username và email.", [io("display_name","String?","≤80 ký tự"),io("username","String","3–50 ký tự"),io("email","Email","≤254 ký tự")], [io("user","UserOut","5 trường công khai")], "1) Kiểm tra quyền sở hữu. 2) Chuẩn hóa và kiểm tra trùng. 3) Vô hiệu token reset nếu email đổi. 4) Commit và trả hồ sơ.", "Đã đăng nhập; User tồn tại.", "Dữ liệu mới được lưu duy nhất; token reset cũ bị xóa khi cần."),
            method("forgot_password", "Tạo yêu cầu khôi phục mật khẩu an toàn.", [io("email","Email","≤254 ký tự")], [io("message","String","Thông báo chung")], "1) Chuẩn hóa email. 2) Kiểm tra giới hạn 3 lần/giờ. 3) Nếu User tồn tại, tạo token, lưu hash và hạn 15 phút. 4) Gửi email; luôn trả thông báo chung.", "Người dùng truy cập chức năng quên mật khẩu.", "Không tiết lộ email tồn tại; token hợp lệ được lưu khi tìm thấy tài khoản."),
            method("reset_password", "Đổi mật khẩu bằng token dùng một lần.", [io("token","String","Token URL-safe"),io("new_password","String","≥6 ký tự")], [io("message","String","Thông báo kết quả")], "1) Hash token đầu vào. 2) Tìm User và kiểm tra hạn. 3) Băm mật khẩu mới. 4) Xóa token và commit.", "Token tồn tại, chưa hết hạn và chưa sử dụng.", "Mật khẩu mới được lưu; token bị vô hiệu hóa; người dùng đăng nhập lại."),
            method("upload_avatar", "Tải và gắn ảnh đại diện cho User.", [io("file","Binary","≤2 MB"),io("content_type","String","PNG/JPEG/WebP")], [io("user","UserOut","Có avatar_url mới")], "1) Kiểm tra MIME, chữ ký tệp và kích thước. 2) Lưu với tên ngẫu nhiên. 3) Cập nhật avatar_url. 4) Xóa ảnh cũ sau commit.", "Đã đăng nhập; tệp không rỗng.", "Ảnh mới tồn tại và URL được lưu; lỗi thì không thay đổi avatar cũ."),
            method("delete_avatar", "Xóa ảnh đại diện hiện tại.", [], [io("user","UserOut","avatar_url=null")], "1) Lưu URL cũ. 2) Đặt avatar_url=null và commit. 3) Xóa tệp local thuộc thư mục avatar.", "Đã đăng nhập; User tồn tại.", "Hồ sơ không còn avatar_url; tệp local cũ được loại bỏ nếu có."),
        ],
    },
    {
        "name":"Category","table":"categories","desc":"Quản lý danh mục thu/chi thuộc riêng từng người dùng.",
        "attrs":[("id","String/UUID","36 ký tự","Khóa chính."),("user_id","String/UUID","36 ký tự","FK → users.id; bắt buộc."),("name","String","Tên hiển thị","Tên danh mục."),("type","Enum String","income/expense","Loại thu hoặc chi."),("color","Hex String","7 ký tự","Màu #RRGGBB."),("icon","Enum String","Khóa icon cho phép","Biểu tượng giao diện.")],
        "methods":[
            method("list_categories","Liệt kê danh mục của người dùng.", [io("query","String?","Từ khóa tùy chọn")], [io("categories","List<CategoryOut>","0..n phần tử")], "1) Lấy current_user. 2) Lọc Category theo user_id. 3) Lọc tên ở client khi có query. 4) Trả danh sách.", "Đã đăng nhập.", "Chỉ dữ liệu thuộc current_user được trả."),
            method("create_category","Tạo danh mục mới.", [io("name","String","Không rỗng"),io("type","Enum","income/expense"),io("color","Hex String","7 ký tự"),io("icon","Enum String","1 khóa hợp lệ")], [io("category","CategoryOut","6 trường")], "1) Kiểm tra type/icon. 2) Kiểm tra name+type không trùng theo user. 3) Gắn user_id. 4) Commit và trả Category.", "Đã đăng nhập; dữ liệu hợp lệ.", "Danh mục mới thuộc current_user và xuất hiện trong danh sách."),
            method("update_category","Cập nhật tên, màu hoặc icon.", [io("category_id","UUID","36 ký tự"),io("name","String?","Không rỗng nếu có"),io("color","Hex String?","7 ký tự"),io("icon","Enum String?","1 khóa")], [io("category","CategoryOut","6 trường")], "1) Tìm Category theo id+user_id. 2) Kiểm tra dữ liệu và trùng tên. 3) Gán trường được gửi. 4) Commit.", "Đã đăng nhập; Category thuộc current_user.", "Category được cập nhật; các quan hệ được giữ nguyên."),
            method("delete_category","Xóa danh mục không còn được giao dịch sử dụng.", [io("category_id","UUID","36 ký tự")], [io("status","HTTP","204 hoặc lỗi")], "1) Tìm Category theo quyền sở hữu. 2) Kiểm tra Transaction tham chiếu. 3) Nếu không có, xóa và commit; nếu có, trả lỗi.", "Đã đăng nhập; Category tồn tại.", "Thành công: Category bị xóa. Thất bại: dữ liệu giữ nguyên."),
        ],
    },
    {
        "name":"Transaction","table":"transactions","desc":"Lưu khoản thu/chi và cung cấp dữ liệu cho Dashboard, Reports và Assistant.",
        "attrs":[("id","String/UUID","36 ký tự","Khóa chính."),("user_id","String/UUID","36 ký tự","FK → users.id."),("category_id","String/UUID","36 ký tự","FK → categories.id."),("amount","Numeric","14,2","Số tiền chính xác; >0."),("type","Enum String","income/expense","Suy ra từ Category.type."),("txn_date","Date","YYYY-MM-DD","Ngày giao dịch; không ở tương lai."),("note","String?","≤240 ký tự ở UI","Ghi chú tùy chọn.")],
        "methods":[
            method("list_transactions","Liệt kê và lọc giao dịch.", [io("category_id","UUID?","36 ký tự"),io("type","Enum?","income/expense"),io("from_date","Date?","YYYY-MM-DD"),io("to_date","Date?","YYYY-MM-DD"),io("query","String?","Từ khóa")], [io("transactions","List<TransactionOut>","0..n phần tử")], "1) Khởi tạo truy vấn theo user_id. 2) Áp dụng category/type/ngày. 3) Sắp xếp. 4) Client lọc query theo ghi chú hoặc danh mục.", "Đã đăng nhập; khoảng ngày hợp lệ.", "Chỉ giao dịch đúng quyền và điều kiện được trả."),
            method("create_transaction","Tạo khoản thu hoặc chi.", [io("category_id","UUID","36 ký tự"),io("amount","Decimal","14,2; >0"),io("txn_date","Date","Không tương lai"),io("note","String?","≤240 ký tự")], [io("transaction","TransactionOut","7 trường")], "1) Kiểm tra Category thuộc user. 2) Lấy type từ Category. 3) Kiểm tra amount/ngày. 4) Lưu Transaction và trả kết quả.", "Đã đăng nhập; Category tồn tại.", "Transaction được commit với user_id và type chính xác."),
            method("update_transaction","Cập nhật giao dịch hiện có.", [io("transaction_id","UUID","36 ký tự"),io("category_id","UUID?","36 ký tự"),io("amount","Decimal?","14,2; >0"),io("txn_date","Date?","Không tương lai"),io("note","String?","≤240 ký tự")], [io("transaction","TransactionOut","7 trường")], "1) Tìm Transaction theo id+user_id. 2) Kiểm tra Category mới nếu có. 3) Đồng bộ type. 4) Gán trường và commit.", "Đã đăng nhập; Transaction thuộc current_user.", "Dữ liệu mới được lưu; type nhất quán với Category."),
            method("delete_transaction","Xóa giao dịch thuộc người dùng.", [io("transaction_id","UUID","36 ký tự")], [io("status","HTTP","204 hoặc lỗi")], "1) Tìm theo transaction_id+user_id. 2) Nếu tồn tại, xóa. 3) Commit.", "Đã đăng nhập; người dùng xác nhận xóa.", "Transaction không còn tồn tại; số liệu tổng hợp thay đổi theo dữ liệu còn lại."),
        ],
    },
    {
        "name":"Budget","table":"budgets","desc":"Lưu hạn mức chi tiêu theo danh mục, tháng và năm.",
        "attrs":[("id","String/UUID","36 ký tự","Khóa chính."),("user_id","String/UUID","36 ký tự","FK → users.id."),("category_id","String/UUID","36 ký tự","FK → categories.id."),("month","Integer","1–12","Tháng áp dụng."),("year","Integer","4 chữ số","Năm áp dụng."),("limit_amount","Numeric","14,2","Hạn mức; >0.")],
        "methods":[
            method("list_budgets","Lấy ngân sách theo kỳ và tính tiến độ.", [io("month","Integer","1–12"),io("year","Integer","4 chữ số")], [io("budgets","List<BudgetOut>","0..n phần tử")], "1) Lọc Budget theo user và kỳ. 2) SUM giao dịch chi theo category. 3) Tính spent, remaining, usage và is_over. 4) Trả kết quả.", "Đã đăng nhập; kỳ hợp lệ.", "Danh sách gồm dữ liệu lưu và trường tính động chính xác."),
            method("create_budget","Đặt ngân sách mới.", [io("category_id","UUID","36 ký tự"),io("month","Integer","1–12"),io("year","Integer","4 chữ số"),io("limit_amount","Decimal","14,2; >0")], [io("budget","BudgetOut","Có spent/is_over")], "1) Kiểm tra Category expense thuộc user. 2) Kiểm tra chưa có Budget cùng category/kỳ. 3) Lưu. 4) Tính và trả dữ liệu đầu ra.", "Đã đăng nhập; Category chi tồn tại.", "Một Budget duy nhất được tạo cho category/month/year/user."),
            method("update_budget","Thay đổi hạn mức ngân sách.", [io("budget_id","UUID","36 ký tự"),io("limit_amount","Decimal","14,2; >0")], [io("budget","BudgetOut","Có tiến độ mới")], "1) Tìm Budget theo quyền sở hữu. 2) Kiểm tra hạn mức. 3) Cập nhật và commit. 4) Tính lại spent/is_over.", "Đã đăng nhập; Budget thuộc current_user.", "Hạn mức và tỷ lệ sử dụng phản ánh giá trị mới."),
            method("delete_budget","Xóa hạn mức ngân sách.", [io("budget_id","UUID","36 ký tự")], [io("status","HTTP","204 hoặc lỗi")], "1) Tìm Budget theo id+user_id. 2) Xóa. 3) Commit.", "Đã đăng nhập; Budget tồn tại.", "Budget bị xóa; giao dịch và Category không bị ảnh hưởng."),
            method("compute_spent","Tính tổng chi của danh mục trong kỳ.", [io("category_id","UUID","36 ký tự"),io("month","Integer","1–12"),io("year","Integer","4 chữ số")], [io("spent","Decimal","14,2")], "1) Xác định đầu/cuối tháng. 2) Lọc Transaction expense theo user/category/ngày. 3) SUM amount; null chuyển thành 0.", "Category và kỳ đã xác định.", "Trả tổng chi không âm; không ghi thêm dữ liệu."),
        ],
    },
    {
        "name":"SavingGoal","table":"saving_goals","desc":"Quản lý mục tiêu tiết kiệm, tiến độ tính động và kết quả hoàn thành.",
        "attrs":[("id","String/UUID","36 ký tự","Khóa chính."),("user_id","String/UUID","36 ký tự","FK → users.id."),("name","String","1–100 ký tự","Tên mục tiêu."),("target_amount","Numeric","14,2","Số tiền cần đạt; >0."),("deadline","Date?","YYYY-MM-DD","Thời hạn tùy chọn."),("status","Enum String","active/completed","Trạng thái."),("completion_amount","Numeric?","14,2","Số tiền khi hoàn thành."),("completion_mode","Enum String?","keep/release/spend","Cách xử lý tiền.")],
        "methods":[
            method("list_goals","Liệt kê mục tiêu và tiến độ.", [], [io("goals","List<GoalOut>","0..n phần tử")], "1) Lọc SavingGoal theo user_id. 2) Lấy GoalItem. 3) Tính current_amount từ GoalTransaction. 4) Trả danh sách.", "Đã đăng nhập.", "Mỗi GoalOut có items và current_amount tính đúng."),
            method("create_goal","Tạo mục tiêu tiết kiệm.", [io("name","String","1–100 ký tự"),io("target_amount","Decimal","14,2; >0"),io("deadline","Date?","YYYY-MM-DD")], [io("goal","GoalOut","Mục tiêu mới")], "1) Chuẩn hóa tên. 2) Kiểm tra số tiền. 3) Gắn user_id và status=active. 4) Commit và trả GoalOut.", "Đã đăng nhập; dữ liệu hợp lệ.", "Mục tiêu active được tạo với current_amount=0."),
            method("get_goal_detail","Lấy chi tiết mục tiêu.", [io("goal_id","UUID","36 ký tự")], [io("goal","GoalOut","Kèm items/current")], "1) Tìm goal theo id+user_id. 2) Lấy items. 3) Tính current_amount và lịch sử liên quan. 4) Trả chi tiết.", "Đã đăng nhập; goal thuộc current_user.", "Chi tiết đúng quyền được trả hoặc 404."),
            method("update_goal","Cập nhật tên, số tiền đích hoặc thời hạn.", [io("goal_id","UUID","36 ký tự"),io("name","String?","1–100 ký tự"),io("target_amount","Decimal?","14,2; >0"),io("deadline","Date?","YYYY-MM-DD")], [io("goal","GoalOut","Dữ liệu mới")], "1) Kiểm tra quyền và status. 2) Kiểm tra trường gửi lên. 3) Cập nhật. 4) Commit và tái tính tiến độ.", "Đã đăng nhập; goal tồn tại.", "Thông tin mới được lưu; lịch sử tiền không đổi."),
            method("delete_goal","Xóa mục tiêu chưa có lịch sử tiền.", [io("goal_id","UUID","36 ký tự")], [io("status","HTTP","204 hoặc lỗi")], "1) Kiểm tra quyền. 2) Đếm GoalTransaction. 3) Có lịch sử thì chặn. 4) Nếu không, xóa GoalItem rồi SavingGoal và commit.", "Đã đăng nhập; người dùng xác nhận.", "Thành công: goal/items bị xóa. Thất bại: toàn bộ dữ liệu giữ nguyên."),
            method("complete_goal","Hoàn thành và xử lý khoản tiền đã dành.", [io("goal_id","UUID","36 ký tự"),io("mode","Enum","keep/release/spend"),io("category_id","UUID?","36 ký tự khi spend"),io("note","String?","Ghi chú")], [io("goal","GoalOut","status=completed")], "1) Tính current_amount và kiểm tra đạt target. 2) Theo mode: giữ, tạo withdraw, hoặc tạo Transaction chi. 3) Ghi completion_amount/mode. 4) Commit nguyên tử.", "Goal active, thuộc user và current_amount≥target_amount.", "Goal completed; tiền được xử lý đúng mode; lỗi thì rollback."),
        ],
    },
    {
        "name":"GoalItem","table":"goal_items","desc":"Chia mục tiêu tiết kiệm thành các hạng mục kế hoạch nhỏ.",
        "attrs":[("id","String/UUID","36 ký tự","Khóa chính."),("goal_id","String/UUID","36 ký tự","FK → saving_goals.id."),("name","String","Không rỗng","Tên hạng mục."),("cost","Numeric","14,2","Chi phí; >0."),("is_purchased","Boolean","1 bit logic","Trạng thái hoàn tất; mặc định false.")],
        "methods":[
            method("add_item","Thêm hạng mục vào mục tiêu.", [io("goal_id","UUID","36 ký tự"),io("name","String","Không rỗng"),io("cost","Decimal","14,2; >0")], [io("goal","GoalOut","Kèm item mới")], "1) Kiểm tra goal thuộc user. 2) Kiểm tra tên/cost. 3) Tạo GoalItem. 4) Commit và trả GoalOut.", "Đã đăng nhập; SavingGoal tồn tại.", "Item thuộc đúng goal và xuất hiện trong kế hoạch."),
            method("update_item","Cập nhật nội dung hoặc trạng thái hạng mục.", [io("goal_id","UUID","36 ký tự"),io("item_id","UUID","36 ký tự"),io("name","String?","Không rỗng"),io("cost","Decimal?","14,2; >0"),io("is_purchased","Boolean?","true/false")], [io("goal","GoalOut","Kèm item mới")], "1) Kiểm tra quyền goal. 2) Tìm item trong goal. 3) Gán trường hợp lệ. 4) Commit.", "Goal và Item tồn tại, đúng quan hệ.", "Item được cập nhật; tổng cost kế hoạch phản ánh dữ liệu mới."),
            method("delete_item","Xóa hạng mục con.", [io("goal_id","UUID","36 ký tự"),io("item_id","UUID","36 ký tự")], [io("goal","GoalOut","Không còn item")], "1) Kiểm tra quyền goal. 2) Tìm item theo goal_id. 3) Xóa và commit. 4) Trả GoalOut.", "Đã đăng nhập; item thuộc goal.", "Item bị xóa; SavingGoal và lịch sử tiền không thay đổi."),
        ],
    },
    {
        "name":"GoalTransaction","table":"goal_transactions","desc":"Lưu nhật ký nạp/rút và là nguồn tính current_amount của mục tiêu.",
        "attrs":[("id","String/UUID","36 ký tự","Khóa chính."),("goal_id","String/UUID","36 ký tự","FK → saving_goals.id."),("amount","Numeric","14,2","Số tiền; >0."),("type","Enum String","deposit/withdraw","Loại biến động."),("txn_date","Date","YYYY-MM-DD","Ngày ghi nhận."),("note","String?","Ghi chú tùy chọn","Nguồn tiền hoặc lý do.")],
        "methods":[
            method("list_goal_transactions","Lấy lịch sử nạp/rút của mục tiêu.", [io("goal_id","UUID","36 ký tự")], [io("transactions","List<GoalTransactionOut>","0..n phần tử")], "1) Kiểm tra goal thuộc user. 2) Lọc theo goal_id. 3) Sắp xếp mới nhất trước. 4) Trả danh sách.", "Đã đăng nhập; goal tồn tại.", "Chỉ lịch sử của goal đúng quyền được trả."),
            method("deposit","Dành tiền vào mục tiêu.", [io("goal_id","UUID","36 ký tự"),io("amount","Decimal","14,2; >0"),io("note","String?","Ghi chú")], [io("goal","GoalOut","current tăng")], "1) Kiểm tra goal active. 2) Tính available_balance. 3) Chặn nếu amount vượt số dư. 4) Tạo deposit và commit.", "Goal active, thuộc user; amount hợp lệ.", "Deposit được ghi; current_amount tăng đúng amount; lỗi thì không ghi."),
            method("withdraw","Rút tiền khỏi mục tiêu.", [io("goal_id","UUID","36 ký tự"),io("amount","Decimal","14,2; >0"),io("note","String?","Ghi chú")], [io("goal","GoalOut","current giảm")], "1) Kiểm tra goal active. 2) Tính current_amount. 3) Chặn nếu amount vượt current. 4) Tạo withdraw và commit.", "Goal active, thuộc user; current_amount đủ.", "Withdraw được ghi; current_amount giảm; tiền trở lại số dư khả dụng."),
            method("compute_current","Tính số tiền hiện có của mục tiêu.", [io("goal_id","UUID","36 ký tự")], [io("current_amount","Decimal","14,2")], "1) SUM các bản ghi deposit. 2) SUM các bản ghi withdraw. 3) Trả deposit−withdraw.", "Goal tồn tại.", "Kết quả không được lưu thành cột và phản ánh toàn bộ lịch sử."),
        ],
    },
    {
        "name":"AIConversation","table":"ai_conversations","desc":"Đại diện một cuộc hội thoại giữa người dùng và Trợ lý AI.",
        "attrs":[("id","String/UUID","36 ký tự","Khóa chính."),("user_id","String/UUID","36 ký tự","FK → users.id; có index."),("title","String","Tối đa 160 ký tự","Tiêu đề hội thoại."),("created_at","DateTime(TZ)","Ngày giờ + múi giờ","Thời điểm tạo."),("updated_at","DateTime(TZ)","Ngày giờ + múi giờ","Cập nhật khi có trả lời.")],
        "methods":[
            method("list_conversations","Liệt kê hội thoại của người dùng.", [], [io("conversations","List<Summary>","0..n phần tử")], "1) Lọc theo user_id. 2) Lấy preview message. 3) Sắp xếp updated_at giảm dần. 4) Trả danh sách.", "Đã đăng nhập.", "Không lộ hội thoại của người dùng khác."),
            method("get_conversation","Lấy hội thoại và toàn bộ message.", [io("conversation_id","UUID","36 ký tự")], [io("conversation","ConversationOut","Kèm messages")], "1) Tìm conversation theo id+user_id. 2) Lấy messages theo thời gian. 3) Giải mã evidence_json. 4) Trả kết quả.", "Đã đăng nhập; conversation thuộc current_user.", "Hội thoại đúng quyền được trả hoặc 404."),
            method("create_conversation","Tạo hội thoại khi gửi câu hỏi đầu tiên.", [io("question","String","Không rỗng")], [io("conversation","AIConversation","title≤160")], "1) Lấy tối đa 72 ký tự đầu câu hỏi làm title. 2) Gắn user_id. 3) Flush để có id. 4) Chỉ commit cùng messages sau khi AI thành công.", "Đã đăng nhập; không truyền conversation_id.", "Thành công: conversation có cặp message đầu tiên. Lỗi AI: rollback."),
            method("delete_conversation","Xóa hội thoại và messages.", [io("conversation_id","UUID","36 ký tự")], [io("status","HTTP","204 hoặc lỗi")], "1) Kiểm tra quyền sở hữu. 2) Xóa AIMessage theo conversation_id. 3) Xóa AIConversation. 4) Commit.", "Đã đăng nhập; người dùng xác nhận.", "Conversation và messages liên quan không còn tồn tại."),
        ],
    },
    {
        "name":"AIMessage","table":"ai_messages","desc":"Lưu câu hỏi, câu trả lời và snapshot bằng chứng dùng để sinh tư vấn.",
        "attrs":[("id","String/UUID","36 ký tự","Khóa chính."),("conversation_id","String/UUID","36 ký tự","FK → ai_conversations.id; có index."),("role","Enum String","Tối đa 16 ký tự","user hoặc assistant."),("content","Text","Văn bản dài","Nội dung message."),("context_month","Integer?","1–12","Tháng evidence."),("context_year","Integer?","4 chữ số","Năm evidence."),("evidence_json","Text?","JSON văn bản","Snapshot số liệu tổng hợp."),("created_at","DateTime(TZ)","Ngày giờ + múi giờ","Thời điểm tạo.")],
        "methods":[
            method("ask_assistant","Gửi câu hỏi tới Gemini và lưu cặp message.", [io("question","String","Không rỗng"),io("month","Integer?","1–12"),io("year","Integer?","4 chữ số"),io("conversation_id","UUID?","36 ký tự")], [io("reply","AssistantReplyOut","Conversation + model")], "1) Kiểm tra/khởi tạo conversation. 2) Tạo evidence tổng hợp. 3) Gọi Gemini với history. 4) Nếu thành công, lưu user+assistant message và evidence; nếu lỗi, rollback.", "Đã đăng nhập; câu hỏi không rỗng; conversation đúng quyền nếu được truyền.", "Thành công: hai message được commit và updated_at đổi. Lỗi AI: không lưu message mới."),
            method("message_out","Chuyển model AIMessage thành dữ liệu trả về.", [io("message","AIMessage","1 bản ghi")], [io("message_out","AssistantMessageOut","Có evidence object")], "1) Đọc role/content/kỳ. 2) Giải mã evidence_json nếu có. 3) Chuẩn hóa dữ liệu. 4) Trả schema response.", "AIMessage tồn tại và thuộc conversation đúng quyền.", "Dữ liệu đầu ra đọc được; JSON lỗi được xử lý an toàn theo router."),
        ],
    },
]


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr(); shd = tc_pr.find(qn("w:shd"))
    if shd is None: shd = OxmlElement("w:shd"); tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr(); elem = OxmlElement("w:tblHeader"); elem.set(qn("w:val"), "true"); tr_pr.append(elem)


def set_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders"); tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single"); element.set(qn("w:sz"), "4"); element.set(qn("w:color"), "AAB7C4")
        borders.append(element)


def style_table(table, header=True, font_size=8.5):
    set_borders(table)
    if header:
        set_repeat_header(table.rows[0])
        for cell in table.rows[0].cells:
            shade(cell, LIGHT_BLUE)
            for run in cell.paragraphs[0].runs: run.bold = True
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                for run in p.runs:
                    run.font.name = "Arial"; run.font.size = Pt(font_size)


def add_io_table(cell, values):
    if not values:
        cell.paragraphs[0].add_run("Không có.")
        return
    table = cell.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "Tên"
    table.rows[0].cells[1].text = "Kiểu dữ liệu"
    table.rows[0].cells[2].text = "Kích thước"
    for value in values:
        cells = table.add_row().cells
        for i, item in enumerate(value): cells[i].text = item
    style_table(table, True, 8)


def add_method(doc, item, number):
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 3"]
    p.add_run(f"Phương thức {number}: {item['name']}")
    table = doc.add_table(rows=0, cols=2)
    set_borders(table)
    rows = [
        ("Tên", item["name"]),
        ("Mô tả", item["desc"]),
        ("Tham số đầu vào", item["inputs"]),
        ("Kết quả đầu ra", item["outputs"]),
        ("Luồng xử lý", item["flow"]),
        ("Điều kiện bắt đầu", item["pre"]),
        ("Điều kiện kết thúc", item["post"]),
    ]
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label; shade(cells[0], LIGHT_BLUE)
        cells[0].paragraphs[0].runs[0].bold = True
        if label in {"Tham số đầu vào", "Kết quả đầu ra"}: add_io_table(cells[1], value)
        else: cells[1].text = value
    for row in table.rows:
        row.cells[0].width = Inches(1.45); row.cells[1].width = Inches(5.65)
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                for run in p.runs: run.font.name = "Arial"; run.font.size = Pt(8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def remove_part_two(doc):
    target = next(p for p in doc.paragraphs if p.text.strip().startswith("2. Đặc tả các lớp"))
    body = doc._element.body; start = list(body).index(target._p)
    for element in list(body)[start:]:
        if element.tag != qn("w:sectPr"): body.remove(element)


def build():
    doc = Document(SOURCE)
    remove_part_two(doc)
    doc.add_heading("2. Đặc tả Class (Class Specification)", level=1)
    p = doc.add_paragraph(
        "Phần này đặc tả 9 lớp thực thể trong mô hình lớp. Các lớp ORM hiện chủ yếu chứa thuộc tính; "
        "phương thức nghiệp vụ được cài đặt tại Router/Core nhưng được trình bày dưới lớp liên quan để "
        "thể hiện đầy đủ trách nhiệm, dữ liệu vào/ra và điều kiện xử lý."
    )
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for class_index, spec in enumerate(SPECS, start=1):
        if class_index > 1: doc.add_page_break()
        doc.add_heading(f"2.{class_index}. Đặc tả Class {spec['name']}", level=2)
        p = doc.add_paragraph()
        p.add_run("Bảng ánh xạ: ").bold = True; p.add_run(spec["table"])
        p = doc.add_paragraph(spec["desc"]); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        doc.add_heading("a) Các thuộc tính", level=3)
        table = doc.add_table(rows=1, cols=4)
        for i, value in enumerate(("Tên", "Kiểu dữ liệu", "Kích thước", "Ý nghĩa")): table.rows[0].cells[i].text = value
        for attr in spec["attrs"]:
            cells = table.add_row().cells
            for i, value in enumerate(attr): cells[i].text = value
        style_table(table, True, 8.5)
        doc.add_heading("b) Các phương thức", level=3)
        for method_index, item in enumerate(spec["methods"], start=1): add_method(doc, item, method_index)
    doc.core_properties.title = "Đặc tả Class — Sổ Chi Tiêu"
    doc.core_properties.subject = "Phần 2 đặc tả thuộc tính và phương thức theo mã nguồn as-built"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__": build()
