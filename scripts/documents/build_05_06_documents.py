from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parents[2]


ROOT = PROJECT_ROOT
DOWNLOADS = Path(r"C:\Users\LENOVO\Downloads")
SOURCE_05 = DOWNLOADS / "05_GenAI_SoftwareDevelopment_functional-testing.docx"
SOURCE_06 = DOWNLOADS / "06_GenAI_SoftwareDevelopment_screenflow_db.docx"
OUTPUT_05 = (PROJECT_ROOT / "docs/archive/05_GenAI_SoftwareDevelopment_functional-testing_completed.docx")
OUTPUT_06 = (PROJECT_ROOT / "docs/archive/ban-khac/06_GenAI_SoftwareDevelopment_screenflow_db_completed.docx")
DOWNLOAD_05 = DOWNLOADS / OUTPUT_05.name
DOWNLOAD_06 = DOWNLOADS / OUTPUT_06.name
SCREEN_FLOW = (PROJECT_ROOT / ".tools/document-assets/screenflow_db_media/ScreenFlow_AsBuilt.png")
LOGIN_FLOW = (PROJECT_ROOT / ".tools/document-assets/screenflow_db_media/LoginFlow_Visual.png")
ERD = (PROJECT_ROOT / ".tools/document-assets/screenflow_db_media/ERD_AsBuilt.png")
DETAILED_FLOW_DIR = (PROJECT_ROOT / ".tools/document-assets/screenflow_db_media/detailed")
DETAILED_FLOWS = [
    ("01_Authentication_Flow.png", "Luồng xác thực người dùng"),
    ("02_Transactions_Flow.png", "Luồng quản lý giao dịch"),
    ("03_Categories_Flow.png", "Luồng quản lý danh mục"),
    ("04_Budgets_Flow.png", "Luồng quản lý ngân sách"),
    ("05_Goals_Flow.png", "Luồng quản lý mục tiêu tiết kiệm"),
    ("06_Reports_Flow.png", "Luồng báo cáo và đối chiếu"),
    ("07_Assistant_Flow.png", "Luồng Trợ lý AI"),
    ("08_Profile_Flow.png", "Luồng quản lý hồ sơ"),
]
SCREENSHOT_DIR = (PROJECT_ROOT / ".tools/document-assets/screenflow_source_screens/New folder")
SCREENSHOT_STEPS = [
    ("Xác thực người dùng", "Màn hình đăng nhập", "login.png", "Người dùng truy cập /login; có thể nhập thông tin, chuyển sang đăng ký hoặc quên mật khẩu."),
    ("Xác thực người dùng", "Nhập thông tin đăng nhập", "login fill.png", "Người dùng nhập tên đăng nhập/email và mật khẩu, sau đó nhấn Đăng nhập."),
    ("Xác thực người dùng", "Thông báo đăng nhập thất bại", "login error.png", "Thông tin không hợp lệ; hệ thống giữ người dùng tại /login và hiển thị lỗi để thử lại."),
    ("Xác thực người dùng", "Màn hình đăng ký", "register.png", "Người dùng tạo tài khoản; đăng ký thành công sẽ quay về màn hình đăng nhập."),
    ("Xác thực người dùng", "Khôi phục mật khẩu", "reset password.png", "Người dùng nhập email để nhận liên kết đặt lại mật khẩu."),
    ("Tổng quan", "Dashboard sau khi đăng nhập", "dashboard.png", "Đăng nhập thành công; hệ thống chuyển đến /dashboard và hiển thị tổng quan tài chính."),
    ("Quản lý giao dịch", "Danh sách giao dịch", "transactions.png", "Người dùng xem lịch sử, số liệu trong kỳ và các công cụ quản lý giao dịch."),
    ("Quản lý giao dịch", "Tìm kiếm và lọc giao dịch", "Search transactions.png", "Người dùng lọc theo loại, danh mục, thời gian hoặc từ khóa; danh sách được cập nhật theo điều kiện."),
    ("Quản lý giao dịch", "Tạo giao dịch", "Create transaction.png", "Người dùng nhập loại giao dịch, số tiền, danh mục, ngày và ghi chú rồi lưu."),
    ("Quản lý danh mục", "Danh sách danh mục", "Categories.png", "Người dùng xem và quản lý riêng các danh mục thu và danh mục chi."),
    ("Quản lý danh mục", "Tạo danh mục", "Create category.png", "Người dùng chọn loại, nhập tên, chọn biểu tượng và màu rồi lưu danh mục mới."),
    ("Quản lý ngân sách", "Danh sách ngân sách", "budgets.png", "Người dùng theo dõi hạn mức, số đã chi và trạng thái sử dụng ngân sách trong kỳ."),
    ("Quản lý ngân sách", "Đặt ngân sách", "create buget.png", "Người dùng chọn danh mục chi, nhập hạn mức và lưu ngân sách."),
    ("Quản lý mục tiêu", "Danh sách mục tiêu tiết kiệm", "Goal saving.png", "Người dùng theo dõi tiến độ các mục tiêu và mở chức năng tạo mục tiêu mới."),
    ("Quản lý mục tiêu", "Tạo mục tiêu tiết kiệm", "Create goal saving.png", "Người dùng nhập tên, số tiền mục tiêu và thông tin liên quan rồi lưu."),
    ("Quản lý mục tiêu", "Chi tiết mục tiêu tiết kiệm", "Details goal saving.png", "Người dùng xem tiến độ, hạng mục con và lịch sử giao dịch của mục tiêu."),
    ("Quản lý mục tiêu", "Nạp tiền vào mục tiêu", "send money to goal saving.png", "Người dùng nhập số tiền nạp và xác nhận; hệ thống kiểm tra số dư trước khi ghi nhận."),
    ("Báo cáo", "Báo cáo so sánh tài chính", "Reports.png", "Người dùng chọn kỳ và xem so sánh thu, chi, chênh lệch cùng các dữ liệu phân bổ."),
    ("Trợ lý AI", "Màn hình Trợ lý AI", "Ai chatbot.png", "Người dùng chọn hội thoại, kỳ dữ liệu và nhập câu hỏi tài chính."),
    ("Trợ lý AI", "Câu trả lời của Trợ lý AI", "AI answer.png", "Hệ thống hiển thị câu trả lời của Gemini cùng dữ liệu evidence để đối chiếu."),
    ("Hồ sơ cá nhân", "Cập nhật hồ sơ", "Profile.png", "Người dùng mở SidePanel từ avatar để sửa thông tin hoặc cập nhật ảnh đại diện."),
]
BROWSER_FLOW_DIR = (PROJECT_ROOT / ".tools/document-assets/screenflows_v3/diagrams")
BROWSER_FLOWS = [
    ("01_login.png", "Luồng đăng nhập"),
    ("02_register.png", "Luồng đăng ký"),
    ("03_password_recovery.png", "Luồng khôi phục mật khẩu"),
    ("04_dashboard.png", "Luồng Dashboard"),
    ("05_transaction_create.png", "Luồng tạo giao dịch"),
    ("06_transaction_filter.png", "Luồng tìm kiếm giao dịch"),
    ("07_transaction_edit_delete.png", "Luồng sửa và xóa giao dịch"),
    ("08_category_create.png", "Luồng tạo danh mục"),
    ("09_category_edit_delete.png", "Luồng sửa và xóa danh mục"),
    ("10_budget_create.png", "Luồng đặt ngân sách"),
    ("11_budget_edit_delete.png", "Luồng sửa và xóa ngân sách"),
    ("12_goal_create.png", "Luồng tạo mục tiêu"),
    ("13_goal_edit_delete.png", "Luồng sửa và xóa mục tiêu"),
    ("14_goal_items.png", "Luồng quản lý hạng mục con"),
    ("15_goal_deposit.png", "Luồng nạp tiền mục tiêu"),
    ("16_goal_withdraw.png", "Luồng rút tiền mục tiêu"),
    ("17_goal_complete.png", "Luồng hoàn thành mục tiêu"),
    ("18_reports.png", "Luồng báo cáo"),
    ("19_assistant.png", "Luồng Trợ lý AI"),
    ("20_profile.png", "Luồng hồ sơ và ảnh đại diện"),
    ("21_logout.png", "Luồng đăng xuất"),
]


BLUE = "2E74B5"
LIGHT_BLUE = "D9E5F6"
LIGHT_GREEN = "D9F2E2"
LIGHT_YELLOW = "FFF0C9"
LIGHT_PURPLE = "EADFFD"
LIGHT_RED = "FADDDD"


def clear_document(doc):
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def setup_styles(doc):
    def get_or_add(name, style_type, base=None):
        try:
            return doc.styles[name]
        except KeyError:
            style = doc.styles.add_style(name, style_type)
            if base:
                try:
                    style.base_style = doc.styles[base]
                except KeyError:
                    pass
            return style

    normal = get_or_add("Normal", WD_STYLE_TYPE.PARAGRAPH)
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.paragraph_format.space_after = Pt(4)
    for style_name, size, color in [
        ("Title", 24, BLUE),
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 11.5, "1F4D78"),
    ]:
        style = get_or_add(style_name, WD_STYLE_TYPE.PARAGRAPH, "Normal")
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, color="A6A6A6", size="5"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)
        borders.append(element)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def set_cell_text(cell, text, bold=False, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(str(text))
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")


def add_table(doc, headers, rows, size=8.5, header_fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=len(headers))
    try:
        table.style = "Normal Table"
    except KeyError:
        pass
    table.autofit = True
    set_table_borders(table)
    repeat_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for idx, value in enumerate(headers):
        set_cell_text(table.rows[0].cells[idx], value, True, size)
        set_cell_shading(table.rows[0].cells[idx], header_fill)
    for row_values in rows:
        row = table.add_row()
        prevent_row_split(row)
        for idx, value in enumerate(row_values):
            set_cell_text(row.cells[idx], value, False, size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level)
    p.paragraph_format.keep_with_next = True
    return p


def add_text(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.08
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(2)
        p.add_run("• ")
        p.add_run(item)


def add_note(doc, title, text, fill=LIGHT_YELLOW):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Normal Table"
    set_table_borders(table, "D6B656")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(f"{title}: ")
    run.bold = True
    run.font.size = Pt(9.5)
    p.add_run(text).font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_cover(doc, title):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(70)
    run = p.add_run(title)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    run = p.add_run("HỆ THỐNG QUẢN LÝ CHI TIÊU CÁ NHÂN CÓ TÍCH HỢP AI")
    run.bold = True
    run.font.size = Pt(15)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(30)
    p.add_run("Nhóm 21 - Thành viên nhóm 2\n").bold = True
    p.add_run("1. Đào Minh Tâm (Trưởng nhóm)\n2. Phạm Đức Nam")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(25)
    p.add_run("Thời gian thực hiện: 27/07/2026 – 27/09/2026 (9 tuần)")
    doc.add_page_break()


def set_header_footer(section, label):
    header = section.header.paragraphs[0] if section.header.paragraphs else section.header.add_paragraph()
    header.text = label
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(110, 110, 110)
    footer = section.footer.paragraphs[0] if section.footer.paragraphs else section.footer.add_paragraph()
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Nhóm 21  |  Sổ Chi Tiêu  |  Trang ")
    run.font.size = Pt(8)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)


TEST_CASES = [
    ("FT-SYS-01", "Health check", "Kiểm tra API đang chạy.", "Backend khởi động.", "GET /api/health", "HTTP 200; status=ok.", "pytest"),
    ("FT-SYS-02", "Readiness & routes", "Kiểm tra DB và đăng ký route.", "DB kết nối được.", "GET /api/ready", "HTTP 200; status=ready; đủ router.", "pytest"),
    ("FT-AUTH-01", "Đăng ký", "Tạo tài khoản hợp lệ.", "Username/email chưa tồn tại.", "minhtam, tam@example.com, mật khẩu ≥6", "Tạo User; mật khẩu lưu dạng hash.", "pytest + Vitest"),
    ("FT-AUTH-02", "Validation đăng ký", "Thiếu email/username hoặc mật khẩu ngắn.", "Ở trang Đăng ký.", "Email rỗng; username 2 ký tự; password 5 ký tự", "Không gửi hoặc API trả 422/400.", "pytest + Vitest"),
    ("FT-AUTH-03", "Trùng tài khoản", "Kiểm tra username/email duy nhất.", "Đã có tài khoản.", "Email khác hoa/thường; username trùng", "HTTP 400; không tạo bản ghi mới.", "pytest"),
    ("FT-AUTH-04", "Đăng nhập", "Đăng nhập bằng username hoặc email.", "Tài khoản hợp lệ.", "Đúng định danh và mật khẩu", "Trả JWT; tải hồ sơ; vào Dashboard.", "pytest + Vitest"),
    ("FT-AUTH-05", "Sai đăng nhập", "Không tiết lộ thành phần sai.", "Tài khoản tồn tại/không tồn tại.", "Sai mật khẩu hoặc định danh", "HTTP 401; thông báo chung.", "pytest + Vitest"),
    ("FT-AUTH-06", "Khôi phục phiên", "Khôi phục token trong localStorage.", "Có token.", "Token hợp lệ/hết hạn", "Hiển thị AppShell hoặc xóa token và về Login.", "Vitest"),
    ("FT-AUTH-07", "Cập nhật hồ sơ", "Sửa display name, username, email.", "Đã đăng nhập.", "Giá trị hợp lệ", "UserOut mới; AppShell cập nhật ngay.", "pytest + Vitest"),
    ("FT-AUTH-08", "Xung đột hồ sơ", "Dùng email/username của user khác.", "Có hai user.", "Email/username trùng", "HTTP 400; rollback.", "pytest"),
    ("FT-AUTH-09", "Avatar hợp lệ", "Tải lên và xóa avatar.", "Đã đăng nhập.", "PNG/JPEG/WebP ≤2 MB", "Lưu tệp, cập nhật avatar_url; xóa tệp cũ.", "pytest + Vitest"),
    ("FT-AUTH-10", "Avatar không hợp lệ", "Kiểm tra MIME, chữ ký, kích thước.", "Đã đăng nhập.", "Sai MIME/chữ ký, rỗng, >2 MB", "HTTP 400/413/415; không đổi avatar.", "pytest + Vitest"),
    ("FT-AUTH-11", "Quên mật khẩu", "Ẩn việc email có tồn tại và giới hạn tần suất.", "Ở trang Quên mật khẩu.", "Email có/không có; >3 yêu cầu/giờ", "Luôn phản hồi chung; chỉ gửi khi hợp lệ.", "pytest + Vitest"),
    ("FT-AUTH-12", "Đặt lại mật khẩu", "Token đúng, sai, hết hạn và dùng lại.", "Có reset token.", "Token hợp lệ/không hợp lệ", "Token hợp lệ đổi mật khẩu một lần; trường hợp khác HTTP 400.", "pytest + Vitest"),
    ("FT-CAT-01", "Danh sách danh mục", "Chỉ trả dữ liệu của current_user.", "Đã đăng nhập.", "GET /api/categories", "Không lộ danh mục user khác.", "pytest"),
    ("FT-CAT-02", "Tạo danh mục", "Tạo income/expense với màu và icon.", "Đã đăng nhập.", "Ăn uống, expense, #D9A441, utensils", "Tạo Category hợp lệ.", "pytest"),
    ("FT-CAT-03", "Danh mục trùng/icon sai", "Kiểm tra trùng và tập icon.", "Đã có danh mục.", "name+type trùng; icon lạ", "HTTP 400/422.", "pytest"),
    ("FT-CAT-04", "Sửa danh mục", "Cập nhật name/type/color/icon.", "Category thuộc user.", "Dữ liệu mới hợp lệ", "Bản ghi được cập nhật.", "pytest"),
    ("FT-CAT-05", "Xóa danh mục", "Bảo vệ danh mục có giao dịch.", "Category thuộc user.", "Có/không Transaction tham chiếu", "Bị chặn nếu đang dùng; ngược lại xóa.", "pytest"),
    ("FT-TXN-01", "Xem/lọc giao dịch", "Lọc theo category, type, ngày, keyword.", "Có dữ liệu nhiều user.", "Các tổ hợp filter", "Đúng user; sắp ngày giảm dần.", "pytest"),
    ("FT-TXN-02", "Thêm giao dịch", "Tạo giao dịch hợp lệ.", "Category thuộc user.", "amount>0; ngày≤hôm nay", "Tạo; type lấy từ Category.", "pytest"),
    ("FT-TXN-03", "Validation giao dịch", "Chặn tiền không dương/ngày tương lai.", "Đã đăng nhập.", "0, số âm, ngày mai", "HTTP 422; không ghi DB.", "pytest"),
    ("FT-TXN-04", "Quyền danh mục", "Không dùng category của user khác.", "Có hai user.", "category_id của user khác", "HTTP 404.", "pytest"),
    ("FT-TXN-05", "Sửa giao dịch", "Đổi dữ liệu và category.", "Transaction thuộc user.", "category mới hợp lệ", "Cập nhật; đồng bộ type.", "pytest"),
    ("FT-TXN-06", "Xóa giao dịch", "Xóa đúng quyền.", "Có transaction.", "ID đúng/sai/của user khác", "200 hoặc 404; phạm vi user được giữ.", "pytest"),
    ("FT-BUD-01", "Xem ngân sách", "Tính spent và is_over theo kỳ.", "Có Budget và expense.", "month/year", "Tổng chi chính xác; trạng thái vượt đúng.", "pytest"),
    ("FT-BUD-02", "Tạo ngân sách", "Tạo cho danh mục chi.", "Category expense thuộc user.", "limit>0, month 1–12", "Tạo BudgetOut.", "pytest"),
    ("FT-BUD-03", "Ngân sách không hợp lệ", "Chặn trùng kỳ/non-expense.", "Có Budget hoặc Category income.", "Cùng category/month/year", "HTTP 400/404/422.", "pytest"),
    ("FT-BUD-04", "Sửa/xóa ngân sách", "Kiểm tra quyền và hạn mức.", "Budget thuộc user.", "limit mới>0; ID sai", "Cập nhật/xóa hoặc 404/422.", "pytest"),
    ("FT-GOAL-01", "Mục tiêu", "Tạo, xem, sửa mục tiêu.", "Đã đăng nhập.", "Tên≤100; target>0", "GoalOut đúng; scoped theo user.", "pytest"),
    ("FT-GOAL-02", "Hạng mục", "CRUD GoalItem.", "Goal thuộc user.", "name không rỗng; cost>0", "Thêm/sửa/xóa; lỗi 404 khi sai.", "pytest"),
    ("FT-GOAL-03", "Nạp tiền", "Nạp vào goal active.", "Có số dư khả dụng.", "0<amount≤balance", "Tạo deposit; cập nhật current_amount.", "pytest"),
    ("FT-GOAL-04", "Rút tiền", "Rút không vượt tiền đã tiết kiệm.", "Goal có deposit.", "amount≤/>current", "Tạo withdraw hoặc HTTP 400.", "pytest"),
    ("FT-GOAL-05", "Hoàn thành keep", "Hoàn thành và giữ tiền trong goal.", "current≥target.", "mode=keep", "status=completed; không withdraw.", "pytest"),
    ("FT-GOAL-06", "Hoàn thành release", "Trả tiền về số dư khả dụng.", "current≥target.", "mode=release", "Tạo withdraw toàn bộ; lưu completion.", "pytest"),
    ("FT-GOAL-07", "Hoàn thành spend", "Ghi nhận sử dụng tiền mục tiêu.", "current≥target; category expense hợp lệ.", "mode=spend", "Tạo expense + withdraw toàn bộ.", "pytest"),
    ("FT-GOAL-08", "Bảo toàn lịch sử", "Chặn hoàn thành/xóa sai điều kiện.", "Goal chưa đủ/đã completed/có lịch sử.", "Các trạng thái biên", "HTTP 400/409; lịch sử không mất.", "pytest"),
    ("FT-DASH-01", "Dashboard", "Tổng hợp kỳ đã chọn.", "Có dữ liệu tài chính.", "month/year", "Thu, chi, trend, budgets, goals, recent đúng.", "pytest + Vitest"),
    ("FT-DASH-02", "Dashboard rỗng/lỗi", "Empty state và thử lại.", "Kỳ không dữ liệu/API lỗi.", "Kỳ trống; mock lỗi", "Zero/empty state; nút Thử lại hoạt động.", "pytest + Vitest"),
    ("FT-RPT-01", "Báo cáo so sánh", "So sánh kỳ chọn với tháng trước.", "Có giao dịch hai kỳ.", "month/year", "Totals, share, difference, change% đúng.", "pytest + Vitest"),
    ("FT-RPT-02", "Báo cáo rỗng/lỗi", "Empty state và retry.", "Không dữ liệu/API lỗi.", "Kỳ trống; mock lỗi", "has_data=false hoặc hiển thị Thử lại.", "pytest + Vitest"),
    ("FT-AI-01", "Context AI", "Chỉ tạo aggregate an toàn.", "Đã đăng nhập.", "month/year", "Không có identity, note, raw transaction, tên tùy biến.", "pytest"),
    ("FT-AI-02", "Hỏi mới", "Tạo conversation và lưu evidence.", "Gemini mock thành công.", "Câu hỏi hợp lệ", "201; lưu user/assistant messages.", "pytest"),
    ("FT-AI-03", "Hỏi tiếp", "Đưa lịch sử vào prompt.", "Conversation thuộc user.", "conversation_id hợp lệ", "Dùng prior messages; cập nhật updated_at.", "pytest"),
    ("FT-AI-04", "Lịch sử hội thoại", "List/get/delete đúng quyền.", "Có nhiều user/conversation.", "ID đúng/sai", "Không lộ dữ liệu; xóa cả messages.", "pytest"),
    ("FT-AI-05", "Lỗi Gemini", "Map lỗi và rollback.", "Mock service lỗi.", "Thiếu key/429/timeout/upstream", "503/429/504/502; không lưu message.", "pytest + Vitest"),
    ("FT-UI-01", "Auth UI", "Validation và điều hướng form auth.", "Frontend chạy.", "Input hợp lệ/không hợp lệ", "Thông báo đúng; điều hướng đúng.", "Vitest"),
    ("FT-UI-02", "Route guards", "Bảo vệ route và xử lý logout/401.", "Có/không token.", "Route public/protected", "Redirect đúng; token bị xóa khi cần.", "Vitest"),
    ("FT-UI-03", "Profile UI", "Mở panel, lưu profile, kiểm tra avatar.", "Đã đăng nhập.", "Form/profile file", "Cập nhật hoặc hiển thị lỗi tức thời.", "Vitest"),
    ("FT-UI-04", "Retry màn hình", "Dashboard/CRUD/Reports/Assistant xử lý lỗi tải.", "Mock API lỗi.", "Nhấn Thử lại", "Hiển thị lỗi và gọi lại API.", "Vitest"),
    ("FT-UI-05", "Component & API client", "Currency, icon, theme, modal, chuẩn hóa số.", "JSDOM.", "Dữ liệu mẫu", "Render/format/accessibility đúng.", "Vitest"),
    ("FT-BUILD-01", "Production build", "TypeScript compile và Vite build.", "Đã npm ci.", "npm run build", "Sinh frontend/dist không lỗi.", "tsc + Vite"),
]


TEST_REPORT = [
    ("BE-SYSTEM", "02/09/2026", "Đào Minh Tâm, Phạm Đức Nam", "Pass", "—", "3/3 test hệ thống đạt.", "pytest: test_system.py"),
    ("BE-AUTH", "02/09/2026", "Đào Minh Tâm, Phạm Đức Nam", "Pass", "—", "16/16 test xác thực/hồ sơ đạt.", "test_auth*.py"),
    ("BE-CATEGORY", "02/09/2026", "Đào Minh Tâm, Phạm Đức Nam", "Pass", "—", "Icon, scope và CRUD đạt.", "test_category_icons.py + test_crud_routes.py"),
    ("BE-TRANSACTION", "02/09/2026", "Đào Minh Tâm, Phạm Đức Nam", "Pass", "—", "Bộ lọc và mọi nhánh mutation đạt.", "test_crud_routes.py"),
    ("BE-BUDGET", "02/09/2026", "Đào Minh Tâm, Phạm Đức Nam", "Pass", "—", "Spent, filter và CRUD đạt.", "test_crud_routes.py"),
    ("BE-GOAL", "02/09/2026", "Đào Minh Tâm, Phạm Đức Nam", "Pass", "—", "10/10 test mục tiêu đạt.", "test_goals.py + test_goal_branches.py"),
    ("BE-DASHBOARD", "02/09/2026", "Đào Minh Tâm, Phạm Đức Nam", "Pass", "—", "2/2 test tổng hợp/empty đạt.", "test_dashboard.py"),
    ("BE-REPORT", "02/09/2026", "Đào Minh Tâm, Phạm Đức Nam", "Pass", "—", "2/2 test so sánh/empty đạt.", "test_reports.py"),
    ("BE-ASSISTANT", "02/09/2026", "Đào Minh Tâm, Phạm Đức Nam", "Pass", "—", "7/7 test context/hội thoại đạt.", "test_assistant*.py"),
    ("BE-GEMINI", "02/09/2026", "Đào Minh Tâm, Phạm Đức Nam", "Pass", "—", "5/5 test response/retry/timeout đạt.", "test_gemini.py"),
    ("BE-CORE", "02/09/2026", "Đào Minh Tâm, Phạm Đức Nam", "Pass", "—", "18/18 test core/schema/migration đạt.", "test_core_services.py"),
    ("FE-AUTOMATED", "02/09/2026", "Đào Minh Tâm, Phạm Đức Nam", "Pass", "—", "5 file, 50 test frontend đạt.", "Vitest + Testing Library"),
    ("FE-BUILD", "02/09/2026", "Đào Minh Tâm, Phạm Đức Nam", "Pass", "—", "TypeScript/Vite production build đạt.", "1603 modules transformed"),
    ("DB-MIGRATION", "02/09/2026", "Đào Minh Tâm, Phạm Đức Nam", "Pass", "—", "Alembic dựng schema 9 bảng; không có AIReport.", "alembic upgrade/check"),
]


def build_functional_testing():
    # Dựng trên gói OOXML sạch để tránh mang theo thiết lập máy in/header lỗi
    # của tệp mẫu, nhưng vẫn giữ nguyên bố cục và hệ thống mục của mẫu.
    doc = Document()
    setup_styles(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    set_header_footer(section, "KIỂM THỬ CHỨC NĂNG — SỔ CHI TIÊU")

    add_cover(doc, "KIỂM THỬ CHỨC NĂNG ỨNG DỤNG")
    add_heading(doc, "1. Những yêu cầu về tài nguyên cho kiểm thử ứng dụng", 1)
    add_heading(doc, "1.1. Phần cứng", 2)
    add_table(
        doc,
        ["CPU", "RAM", "Ổ đĩa C:", "Architecture", "Kết nối"],
        [("Intel Core i5-12450HX, 8 cores, tối đa 2.4 GHz", "15.7 GB", "174.6 GB; trống 18.8 GB", "64-bit", "Internet cho Gemini/Resend")],
        8.5,
    )
    add_heading(doc, "1.2. Phần mềm", 2)
    add_table(
        doc,
        ["Tên phần mềm", "Phiên bản", "Loại/Mục đích"],
        [
            ("Microsoft Windows 11 Pro", "10.0.26100, 64-bit", "Hệ điều hành kiểm thử"),
            ("Visual Studio Code", "1.135.0", "IDE/Editor"),
            ("Python", "3.14.6", "Backend runtime"),
            ("FastAPI / SQLAlchemy / Pydantic", "0.141.1 / 2.0.52 / 2.13.4", "API, ORM, validation"),
            ("SQLite", "3.53.2", "Database kiểm thử local"),
            ("Node.js / npm", "24.8.0 / 11.19.0", "Frontend toolchain"),
            ("React / TypeScript / Vite", "18.3.1 / 5.6.3 / 8.2.2", "Frontend"),
            ("pytest", "9.1.1", "Kiểm thử backend"),
            ("Vitest / Testing Library", "4.1.11 / 16.3.0", "Kiểm thử frontend"),
            ("Microsoft Edge", "152.0.4191.53", "Trình duyệt kiểm tra giao diện"),
            ("Git", "2.53.0.windows.3", "Quản lý phiên bản"),
        ],
        8.5,
    )
    add_note(
        doc,
        "Phạm vi",
        "Kiểm thử tập trung vào hành vi chức năng của API, validation, quyền sở hữu dữ liệu, "
        "logic tài chính, luồng giao diện và tích hợp Gemini/Resend qua mock. Không gọi dịch vụ "
        "ngoài thật trong bộ test tự động.",
        LIGHT_GREEN,
    )

    landscape = doc.add_section(WD_SECTION.NEW_PAGE)
    landscape.orientation = WD_ORIENT.LANDSCAPE
    landscape.page_width, landscape.page_height = landscape.page_height, landscape.page_width
    landscape.top_margin = Inches(0.45)
    landscape.bottom_margin = Inches(0.45)
    landscape.left_margin = Inches(0.45)
    landscape.right_margin = Inches(0.45)
    set_header_footer(landscape, "KIỂM THỬ CHỨC NĂNG — DANH SÁCH TEST CASE")

    add_heading(doc, "2. Danh sách các tình huống kiểm tra ứng dụng", 1)
    add_text(
        doc,
        f"Tổng cộng {len(TEST_CASES)} tình huống chức năng, bao phủ 11 use case và các nhánh lỗi quan trọng."
    )
    add_table(
        doc,
        ["Test ID", "Chức năng", "Mô tả", "Điều kiện trước", "Dữ liệu test", "Kết quả mong muốn", "Ghi chú"],
        TEST_CASES,
        7.2,
    )

    doc.add_page_break()
    add_heading(doc, "3. Báo cáo kết quả test (Test Report)", 1)
    add_table(
        doc,
        ["Test ID/Suite", "Ngày testing", "Người tham gia test", "Pass/Fail", "Độ nghiêm trọng", "Tóm tắt lỗi/kết quả", "Ghi chú"],
        TEST_REPORT,
        7.5,
    )
    add_heading(doc, "3.1. Tổng hợp kết quả", 2)
    add_table(
        doc,
        ["Hạng mục", "Kết quả", "Đánh giá"],
        [
            ("Backend pytest", "69 passed; 22 subtests passed", "Đạt"),
            ("Frontend Vitest", "5 test files; 50 tests passed", "Đạt"),
            ("Frontend production build", "1603 modules transformed; build thành công", "Đạt"),
            ("Alembic/database", "Schema 9 bảng; integrity check ok; không có AIReport", "Đạt"),
            ("Lỗi chức năng phát hiện", "0 trong phạm vi tự động", "Không có lỗi mở"),
        ],
        8.5,
        LIGHT_GREEN,
    )
    add_note(
        doc,
        "Cảnh báo không chặn",
        "pytest ghi nhận 7 cảnh báo Pydantic V2 về class-based Config. Đây là cảnh báo tương thích "
        "tương lai, không làm test thất bại và không ảnh hưởng kết quả chức năng hiện tại.",
        LIGHT_YELLOW,
    )

    doc.core_properties.title = "Kiểm thử chức năng — Sổ Chi Tiêu"
    doc.core_properties.subject = "Functional testing as-built 02/09/2026"
    doc.core_properties.author = "Nhóm 21"
    doc.save(OUTPUT_05)
    copy2(OUTPUT_05, DOWNLOAD_05)


def build_screenflow_db():
    doc = Document()
    setup_styles(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    set_header_footer(section, "SCREEN FLOW & DATABASE — SỔ CHI TIÊU")

    add_cover(doc, "SCREEN FLOW & TÀI LIỆU THIẾT KẾ CƠ SỞ DỮ LIỆU")

    add_heading(doc, "1. Screen Flow: Phân luồng màn hình của ứng dụng", 1)
    add_heading(doc, "1.1. Phạm vi và nguyên tắc điều hướng", 2)
    add_text(
        doc,
        "Ứng dụng có hai vùng: vùng công khai cho xác thực và vùng được bảo vệ trong AppShell. "
        "Web responsive/PWA và Windows/Tauri sử dụng cùng tập route React và cùng backend FastAPI."
    )
    add_bullets(
        doc,
        [
            "PublicOnlyRoute chuyển người đã đăng nhập khỏi Login/Register/Forgot Password về Dashboard.",
            "ProtectedRoute chuyển người chưa xác thực về Login.",
            "Route không tồn tại được điều hướng về Dashboard.",
            "Profile/Avatar mở bằng SidePanel và không làm mất màn hình nghiệp vụ hiện tại.",
            "Reports và Assistant cung cấp đường dẫn về Transactions để kiểm chứng dữ liệu.",
        ],
    )
    add_heading(doc, "1.2. Screen Flow chi tiết theo chức năng", 2)
    add_text(
        doc,
        "Các sơ đồ sử dụng ảnh chụp trực tiếp từ ứng dụng chạy local với tài khoản và dữ liệu minh họa biệt lập. "
        "Mỗi khung thể hiện một trạng thái giao diện; mũi tên và nhãn mô tả hành động, điều kiện, kết quả "
        "thành công, lỗi hoặc đường quay lại theo mã nguồn hiện hành.",
    )
    for figure_number, (flow_file, flow_title) in enumerate(BROWSER_FLOWS, start=1):
        if figure_number > 1:
            doc.add_page_break()
        add_heading(doc, f"1.2.{figure_number}. {flow_title}", 3)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(BROWSER_FLOW_DIR / flow_file), width=Inches(7.0))
        p = doc.add_paragraph(f"Hình {figure_number}. {flow_title} — ảnh giao diện chụp từ trình duyệt")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.runs[0].italic = True
        p.runs[0].font.size = Pt(9)

    doc.add_page_break()
    add_heading(doc, "1.3. Sơ đồ Screen Flow tổng thể", 2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(SCREEN_FLOW), width=Inches(7.0))
    p = doc.add_paragraph("Hình 22. Screen Flow tổng thể as-built của ứng dụng")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True
    p.runs[0].font.size = Pt(9)

    # Tách danh mục sang trang riêng để không để lại một hàng tiêu đề bảng
    # đơn độc ở cuối trang sơ đồ.
    doc.add_page_break()
    add_heading(doc, "1.4. Danh mục màn hình", 2)
    add_table(
        doc,
        ["Màn hình/Thành phần", "Route", "Quyền truy cập", "Luồng chính"],
        [
            ("Đăng nhập", "/login", "Công khai", "Đăng nhập; sang Đăng ký/Quên mật khẩu."),
            ("Đăng ký", "/register", "Công khai", "Đăng ký thành công → Đăng nhập."),
            ("Quên mật khẩu", "/forgot-password", "Công khai", "Nhập email; nhận liên kết."),
            ("Đặt lại mật khẩu", "/reset-password?token=...", "Công khai", "Đổi mật khẩu → Đăng nhập."),
            ("Dashboard", "/dashboard", "Đã xác thực", "Màn hình mặc định; tổng quan theo kỳ."),
            ("Giao dịch", "/transactions", "Đã xác thực", "CRUD, lọc và nhận deep-link từ Reports/Assistant."),
            ("Danh mục", "/categories", "Đã xác thực", "CRUD danh mục thu/chi."),
            ("Ngân sách", "/budgets", "Đã xác thực", "CRUD ngân sách theo tháng/năm."),
            ("Mục tiêu", "/goals", "Đã xác thực", "Goal/item, nạp/rút, hoàn thành."),
            ("Báo cáo", "/reports", "Đã xác thực", "So sánh kỳ chọn với tháng trước."),
            ("Trợ lý AI", "/assistant", "Đã xác thực", "Evidence, hỏi đáp Gemini, lịch sử."),
            ("Hồ sơ/Avatar", "SidePanel trong AppShell", "Đã xác thực", "Cập nhật hồ sơ, tải/xóa avatar."),
        ],
        8.0,
    )

    doc.add_page_break()
    add_heading(doc, "2. Cơ sở dữ liệu", 1)
    add_heading(doc, "2.1. Cơ sở dữ liệu quan hệ", 2)
    add_text(
        doc,
        "Hệ thống sử dụng SQLAlchemy ORM. Mặc định development dùng SQLite; PostgreSQL được hỗ trợ "
        "qua DATABASE_URL. Schema triển khai được quản lý bằng Alembic revision 20260901_0001 và gồm 9 bảng."
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ERD), width=Inches(7.0))
    p = doc.add_paragraph("Hình 23. ERD schema chốt triển khai — không có bảng AIReport")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True
    p.runs[0].font.size = Pt(9)

    add_heading(doc, "2.1.1. Danh sách bảng", 3)
    add_table(
        doc,
        ["Bảng", "Khóa chính", "Khóa ngoại", "Vai trò"],
        [
            ("users", "id", "—", "Tài khoản, hồ sơ, avatar, reset password."),
            ("categories", "id", "user_id → users", "Danh mục income/expense."),
            ("transactions", "id", "user_id → users; category_id → categories", "Khoản thu/chi."),
            ("budgets", "id", "user_id → users; category_id → categories", "Hạn mức theo danh mục/kỳ."),
            ("saving_goals", "id", "user_id → users", "Mục tiêu và kết quả hoàn thành."),
            ("goal_items", "id", "goal_id → saving_goals", "Hạng mục con."),
            ("goal_transactions", "id", "goal_id → saving_goals", "Nhật ký deposit/withdraw."),
            ("ai_conversations", "id", "user_id → users", "Hội thoại AI."),
            ("ai_messages", "id", "conversation_id → ai_conversations", "Messages và evidence snapshot."),
        ],
        8.0,
    )
    add_heading(doc, "2.1.2. Quan hệ chính", 3)
    add_table(
        doc,
        ["Quan hệ", "Bội số", "Ý nghĩa"],
        [
            ("users – categories/transactions/budgets/saving_goals/ai_conversations", "1 : 0..*", "Một user sở hữu nhiều dữ liệu; mọi truy vấn phải scoped."),
            ("categories – transactions", "1 : 0..*", "Một giao dịch thuộc đúng một danh mục."),
            ("categories – budgets", "1 : 0..*", "Một ngân sách áp dụng cho một danh mục."),
            ("saving_goals – goal_items", "1 : 0..*", "Các hạng mục thuộc mục tiêu."),
            ("saving_goals – goal_transactions", "1 : 0..*", "Nhật ký tiền của mục tiêu."),
            ("ai_conversations – ai_messages", "1 : 0..*", "Các message thuộc một hội thoại."),
        ],
        8.0,
    )

    add_heading(doc, "2.2. Các ràng buộc toàn vẹn trong CSDL", 2)
    add_heading(doc, "2.2.1. Ràng buộc ở database", 3)
    add_table(
        doc,
        ["Nhóm", "Ràng buộc hiện hành"],
        [
            ("Primary key", "Tất cả 9 bảng dùng id kiểu String/UUID."),
            ("Unique/index", "users.username và users.email unique + index; AI conversation/message có index FK."),
            ("Foreign key", "10 quan hệ FK như thể hiện trên ERD; không khai báo ondelete/cascade."),
            ("Not null", "Áp dụng cho trường nhận dạng, FK và dữ liệu nghiệp vụ bắt buộc."),
            ("Kiểu tiền", "Numeric(14,2), không dùng Float."),
            ("Timestamp", "created_at/updated_at dùng DateTime timezone khi phù hợp."),
        ],
        8.5,
    )
    add_heading(doc, "2.2.2. Ràng buộc ở Schema/Router", 3)
    add_table(
        doc,
        ["Đối tượng", "Ràng buộc"],
        [
            ("User", "Username 3–50; password ≥6; email hợp lệ/duy nhất; token reset hạn 15 phút."),
            ("Category", "type ∈ {income, expense}; icon trong danh sách cho phép; name+type không trùng theo user."),
            ("Transaction", "amount>0; ngày không tương lai; category thuộc user; type suy ra từ category."),
            ("Budget", "month 1–12; limit>0; chỉ category expense; không trùng category/month/year theo user."),
            ("SavingGoal", "Tên không rỗng/≤100; target>0; hoàn thành khi current≥target."),
            ("GoalTransaction", "deposit không vượt available balance; withdraw không vượt current; bảo toàn lịch sử."),
            ("AIConversation", "Mọi thao tác kiểm tra user_id; xóa messages trước conversation."),
            ("AIMessage", "role user/assistant; chỉ lưu sau khi Gemini thành công; lỗi thì rollback."),
        ],
        8.0,
    )
    add_heading(doc, "2.2.3. Dữ liệu tính toán không lưu cột", 3)
    add_bullets(
        doc,
        [
            "Budget.spent và is_over được tính từ Transaction theo kỳ.",
            "SavingGoal.current_amount = tổng deposit − tổng withdraw.",
            "Dashboard và Reports được tổng hợp trực tiếp từ dữ liệu giao dịch.",
            "Assistant evidence được tạo từ aggregate an toàn; snapshot chỉ lưu ở AIMessage.",
        ],
    )
    add_note(
        doc,
        "Điểm cần lưu ý khi chốt production",
        "Các ràng buộc tổ hợp Category(name,type,user) và Budget(category,month,year,user) hiện được "
        "kiểm tra ở Router, chưa có UniqueConstraint tại database. Foreign key cũng chưa khai báo "
        "cascade; vòng đời con được điều phối trong nghiệp vụ.",
        LIGHT_YELLOW,
    )

    doc.core_properties.title = "Screen Flow & Database Design — Sổ Chi Tiêu"
    doc.core_properties.subject = "Screen flow và ERD as-built 02/09/2026"
    doc.core_properties.author = "Nhóm 21"
    doc.save(OUTPUT_06)
    copy2(OUTPUT_06, DOWNLOAD_06)


if __name__ == "__main__":
    for required in (SOURCE_05, SOURCE_06, SCREEN_FLOW, ERD, *(BROWSER_FLOW_DIR / item[0] for item in BROWSER_FLOWS)):
        if not required.exists():
            raise FileNotFoundError(required)
    build_functional_testing()
    build_screenflow_db()
    print(OUTPUT_05)
    print(OUTPUT_06)
    print(DOWNLOAD_05)
    print(DOWNLOAD_06)
