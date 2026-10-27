from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parents[2]


ROOT = PROJECT_ROOT
SOURCE = Path(
    r"C:\Users\LENOVO\Downloads\Tai lieu du an"
    r"\04_GenAI_SoftwareDevelopment_object-oriented-design_presentation.docx"
)
OUTPUT = (PROJECT_ROOT / "docs/archive/ban-khac/04_GenAI_SoftwareDevelopment_object-oriented-design_presentation_updated_no_AIReport.docx")
DOWNLOAD_OUTPUT = SOURCE.with_name(
    "04_GenAI_SoftwareDevelopment_object-oriented-design_presentation_updated_no_AIReport.docx"
)
CLASS_DIAGRAM = (PROJECT_ROOT / ".tools/document-assets/ood_updated_media/OOD_Class_Diagram_AsBuilt.png")
ARCH_DIAGRAM = (PROJECT_ROOT / ".tools/document-assets/ood_updated_media/OOD_Architecture_AsBuilt.png")


BLUE = "2E74B5"
LIGHT_BLUE = "D9E5F6"
LIGHT_GREEN = "D9F2E2"
LIGHT_YELLOW = "FFF0C9"
LIGHT_PURPLE = "EADFFD"
LIGHT_RED = "FADDDD"
GRAY = "F2F2F2"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False, size: float = 9.0) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color: str = "A6A6A6", size: str = "6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def add_table(doc, headers, rows, header_fill=LIGHT_BLUE, font_size=9.0):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Normal Table"
    set_table_borders(table)
    table.autofit = True
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, heading in enumerate(headers):
        set_cell_text(header.cells[index], heading, bold=True, size=font_size)
        set_cell_shading(header.cells[index], header_fill)
    for row_values in rows:
        row = table.add_row()
        for index, value in enumerate(row_values):
            set_cell_text(row.cells[index], value, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.paragraph_format.space_after = Pt(2)
        p.add_run("• ")
        p.add_run(item)


def add_numbered(doc, items):
    for index, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.24)
        p.paragraph_format.space_after = Pt(2)
        p.add_run(f"{index}. ")
        p.add_run(item)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(89, 89, 89)


def add_note(doc, title, text, fill=LIGHT_YELLOW):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Normal Table"
    set_table_borders(table, color="D6B656")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{title}: ")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r2 = p.add_run(text)
    r2.font.name = "Arial"
    r2.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_heading(doc, text, level):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.08
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p


CLASS_SPECS = [
    {
        "title": "2.1. Đặc tả lớp User",
        "summary": "Ánh xạ bảng users; lưu danh tính, thông tin hồ sơ, mật khẩu băm và dữ liệu đặt lại mật khẩu.",
        "attrs": [
            ("id", "String (UUID)", "PK; sinh UUID"),
            ("username", "String", "Bắt buộc; duy nhất; có index"),
            ("email", "String", "Bắt buộc; duy nhất; có index"),
            ("display_name", "String?", "Tên hiển thị tùy chọn"),
            ("password_hash", "String", "Mật khẩu đã băm bằng bcrypt"),
            ("avatar_url", "String?", "Đường dẫn ảnh đại diện"),
            ("created_at", "DateTime(timezone)", "Thời điểm tạo"),
            ("reset_token_hash", "String?", "SHA-256 của token đặt lại"),
            ("reset_token_expiry", "DateTime(timezone)?", "Hạn hiệu lực token"),
        ],
        "ops": [
            ("register, login, me", "AuthRouter", "Đăng ký, đăng nhập và lấy người dùng hiện tại"),
            ("update_profile", "AuthRouter", "Cập nhật display_name, username, email"),
            ("upload_avatar, delete_avatar", "AuthRouter + file storage", "Quản lý avatar PNG/JPEG/WebP tối đa 2 MB"),
            ("forgot_password, reset_password", "AuthRouter + Security + Resend", "Khôi phục mật khẩu bằng token dùng một lần"),
        ],
        "relations": "1:N với Category, Transaction, Budget, SavingGoal và AIConversation.",
        "rules": [
            "Username dài 3–50 ký tự; email hợp lệ; username và email duy nhất.",
            "Mật khẩu đầu vào tối thiểu 6 ký tự và chỉ lưu dạng bcrypt hash.",
            "JWT xác thực chứa user id; nghiệp vụ chính lọc dữ liệu theo current_user.id.",
            "Yêu cầu quên mật khẩu tối đa 3 lần/email/giờ trong bộ nhớ; token hết hạn sau 15 phút.",
            "Khi email thay đổi, reset token hiện có bị vô hiệu hóa.",
        ],
    },
    {
        "title": "2.2. Đặc tả lớp Category",
        "summary": "Ánh xạ bảng categories; mô tả danh mục thu/chi riêng của từng người dùng.",
        "attrs": [
            ("id", "String (UUID)", "PK"),
            ("user_id", "String", "FK → users.id"),
            ("name", "String", "Tên danh mục"),
            ("type", "String", "income hoặc expense"),
            ("color", "String", "Màu hiển thị; mặc định #D9A441"),
            ("icon", "String", "Khóa icon; mặc định circle-dollar-sign"),
        ],
        "ops": [
            ("list_categories", "CategoriesRouter", "Liệt kê theo người dùng"),
            ("create_category", "CategoriesRouter", "Tạo danh mục"),
            ("update_category", "CategoriesRouter", "Cập nhật tên, loại, màu, icon"),
            ("delete_category", "CategoriesRouter", "Xóa khi không có giao dịch tham chiếu"),
        ],
        "relations": "N:1 với User; 1:N với Transaction và Budget.",
        "rules": [
            "Cặp name + type không được trùng trong cùng tài khoản.",
            "Type chỉ nhận income hoặc expense; icon phải thuộc tập khóa được Schema cho phép.",
            "Router hiện chặn xóa khi có Transaction tham chiếu; khóa ngoại Budget vẫn phải được bảo toàn.",
        ],
    },
    {
        "title": "2.3. Đặc tả lớp Transaction",
        "summary": "Ánh xạ bảng transactions; lưu khoản thu/chi và là nguồn dữ liệu cho Dashboard, Reports và Assistant.",
        "attrs": [
            ("id", "String (UUID)", "PK"),
            ("user_id", "String", "FK → users.id"),
            ("category_id", "String", "FK → categories.id"),
            ("amount", "Numeric(14,2)", "Số tiền chính xác"),
            ("type", "String", "Suy ra từ Category.type"),
            ("txn_date", "Date", "Ngày giao dịch"),
            ("note", "String?", "Ghi chú tùy chọn"),
        ],
        "ops": [
            ("list_transactions", "TransactionsRouter", "Lọc theo category, type, ngày và từ khóa"),
            ("create_transaction", "TransactionsRouter", "Tạo giao dịch"),
            ("update_transaction", "TransactionsRouter", "Cập nhật và đồng bộ type khi đổi category"),
            ("delete_transaction", "TransactionsRouter", "Xóa giao dịch thuộc người dùng"),
        ],
        "relations": "N:1 với User và Category.",
        "rules": [
            "Amount phải lớn hơn 0; ngày giao dịch không được ở tương lai.",
            "Category phải thuộc người dùng hiện tại.",
            "Client không quyết định type; backend lấy type từ Category.",
        ],
    },
    {
        "title": "2.4. Đặc tả lớp Budget",
        "summary": "Ánh xạ bảng budgets; lưu hạn mức theo danh mục chi và kỳ tháng/năm.",
        "attrs": [
            ("id", "String (UUID)", "PK"),
            ("user_id", "String", "FK → users.id"),
            ("category_id", "String", "FK → categories.id"),
            ("month", "Integer", "1–12"),
            ("year", "Integer", "Năm áp dụng"),
            ("limit_amount", "Numeric(14,2)", "Hạn mức > 0"),
        ],
        "ops": [
            ("list_budgets", "BudgetsRouter", "Liệt kê và tính spent/is_over"),
            ("create_budget", "BudgetsRouter", "Tạo ngân sách"),
            ("update_budget", "BudgetsRouter", "Cập nhật hạn mức"),
            ("delete_budget", "BudgetsRouter", "Xóa ngân sách"),
            ("compute_spent", "BudgetsRouter", "SUM giao dịch chi theo category và kỳ"),
        ],
        "relations": "N:1 với User và Category.",
        "rules": [
            "Chỉ Category loại expense của đúng người dùng được dùng để tạo Budget.",
            "Một category chỉ có một Budget trong cùng month/year theo kiểm tra nghiệp vụ.",
            "spent và is_over là dữ liệu tính động, không phải cột lưu trữ.",
        ],
    },
    {
        "title": "2.5. Đặc tả lớp SavingGoal",
        "summary": "Ánh xạ bảng saving_goals; lưu mục tiêu và kết quả hoàn thành, không lưu current_amount.",
        "attrs": [
            ("id", "String (UUID)", "PK"),
            ("user_id", "String", "FK → users.id"),
            ("name", "String", "Tên mục tiêu"),
            ("target_amount", "Numeric(14,2)", "Mức cần đạt"),
            ("deadline", "Date?", "Hạn tùy chọn"),
            ("status", "String", "active hoặc completed"),
            ("completion_amount", "Numeric(14,2)?", "Số tiền tại lúc hoàn thành"),
            ("completion_mode", "String?", "keep, release hoặc spend"),
        ],
        "ops": [
            ("list/create/update/delete_goal", "GoalsRouter", "Quản lý mục tiêu"),
            ("to_goal_out", "GoalsRouter", "Ghép item và current_amount tính động"),
            ("complete_goal", "GoalsRouter", "Hoàn thành theo keep/release/spend"),
        ],
        "relations": "N:1 với User; 1:N với GoalItem và GoalTransaction.",
        "rules": [
            "Tên không rỗng, tối đa 100 ký tự; target_amount > 0.",
            "current_amount = tổng deposit − tổng withdraw.",
            "Chỉ hoàn thành khi current_amount ≥ target_amount.",
            "Không được xóa mục tiêu đã có lịch sử tiền; GoalItem được xóa thủ công khi xóa mục tiêu hợp lệ.",
        ],
    },
    {
        "title": "2.6. Đặc tả lớp GoalItem",
        "summary": "Ánh xạ bảng goal_items; chia mục tiêu thành các hạng mục nhỏ.",
        "attrs": [
            ("id", "String (UUID)", "PK"),
            ("goal_id", "String", "FK → saving_goals.id"),
            ("name", "String", "Tên hạng mục"),
            ("cost", "Numeric(14,2)", "Chi phí > 0"),
            ("is_purchased", "Boolean", "Mặc định false"),
        ],
        "ops": [
            ("add_item", "GoalsRouter", "Thêm hạng mục"),
            ("update_item", "GoalsRouter", "Sửa tên, cost, trạng thái mua"),
            ("delete_item", "GoalsRouter", "Xóa hạng mục"),
        ],
        "relations": "N:1 với SavingGoal; vòng đời được điều phối qua GoalsRouter.",
        "rules": [
            "Tên không được rỗng; cost phải lớn hơn 0.",
            "Mọi thao tác đều kiểm tra SavingGoal thuộc current_user.",
        ],
    },
    {
        "title": "2.7. Đặc tả lớp GoalTransaction",
        "summary": "Ánh xạ bảng goal_transactions; lưu nhật ký nạp/rút và là cơ sở tính số tiền hiện có.",
        "attrs": [
            ("id", "String (UUID)", "PK"),
            ("goal_id", "String", "FK → saving_goals.id"),
            ("amount", "Numeric(14,2)", "Số tiền > 0"),
            ("type", "String", "deposit hoặc withdraw"),
            ("txn_date", "Date", "Ngày ghi nhận"),
            ("note", "String?", "Ghi chú tùy chọn"),
        ],
        "ops": [
            ("list_goal_transactions", "GoalsRouter", "Lịch sử giảm dần theo ngày"),
            ("deposit", "GoalsRouter + Finance", "Nạp khi goal active và đủ số dư"),
            ("withdraw", "GoalsRouter", "Rút không vượt current_amount"),
            ("compute_current", "GoalsRouter", "SUM deposit − SUM withdraw"),
        ],
        "relations": "N:1 với SavingGoal.",
        "rules": [
            "Amount phải lớn hơn 0.",
            "Nạp không vượt số dư khả dụng; rút không vượt số đã tiết kiệm.",
            "Nhật ký tiền được bảo toàn: mục tiêu có bản ghi này không được xóa.",
        ],
    },
    {
        "title": "2.8. Đặc tả lớp AIConversation",
        "summary": "Ánh xạ bảng ai_conversations; đại diện một luồng hội thoại của người dùng với trợ lý.",
        "attrs": [
            ("id", "String (UUID)", "PK"),
            ("user_id", "String", "FK → users.id; có index"),
            ("title", "String(160)", "Mặc định Cuộc trò chuyện mới"),
            ("created_at", "DateTime(timezone)", "Thời điểm tạo"),
            ("updated_at", "DateTime(timezone)", "Cập nhật khi có câu trả lời"),
        ],
        "ops": [
            ("list_conversations", "AssistantRouter", "Danh sách theo updated_at giảm dần"),
            ("get_conversation", "AssistantRouter", "Lấy hội thoại và messages"),
            ("delete_conversation", "AssistantRouter", "Xóa messages rồi xóa conversation"),
            ("ask_assistant", "AssistantRouter", "Tạo conversation mới khi chưa có id"),
        ],
        "relations": "N:1 với User; 1:N composition với AIMessage ở tầng nghiệp vụ.",
        "rules": [
            "Mọi thao tác kiểm tra conversation_id cùng user_id.",
            "Tiêu đề hội thoại mới lấy tối đa 72 ký tự đầu của câu hỏi.",
        ],
    },
    {
        "title": "2.9. Đặc tả lớp AIMessage",
        "summary": "Ánh xạ bảng ai_messages; lưu câu hỏi, câu trả lời và bằng chứng tổng hợp tại thời điểm trả lời.",
        "attrs": [
            ("id", "String (UUID)", "PK"),
            ("conversation_id", "String", "FK → ai_conversations.id; có index"),
            ("role", "String(16)", "user hoặc assistant"),
            ("content", "Text", "Nội dung thông điệp"),
            ("context_month", "Integer?", "Tháng của evidence"),
            ("context_year", "Integer?", "Năm của evidence"),
            ("evidence_json", "Text?", "Snapshot dữ liệu tổng hợp"),
            ("created_at", "DateTime(timezone)", "Thời điểm tạo"),
        ],
        "ops": [
            ("_message_out", "AssistantRouter", "Giải mã evidence_json cho response"),
            ("ask_assistant", "AssistantRouter", "Lưu cặp user/assistant sau khi Gemini thành công"),
        ],
        "relations": "N:1 với AIConversation.",
        "rules": [
            "Role đầu ra chỉ là user hoặc assistant.",
            "Không lưu message mới khi Gemini lỗi; phiên DB được rollback.",
            "Evidence gắn kỳ và số liệu tổng hợp để có thể giải thích câu trả lời.",
        ],
    },
]


def build_document():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if not CLASS_DIAGRAM.exists() or not ARCH_DIAGRAM.exists():
        raise FileNotFoundError("Thiếu ảnh sơ đồ đã cập nhật.")

    doc = Document(SOURCE)
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)

    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.paragraph_format.space_after = Pt(5)

    for style_name, size, color in [
        ("Title", 26, BLUE),
        ("Heading 1", 16, BLUE),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 11.5, "1F4D78"),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

    # Cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(70)
    r = p.add_run("TÀI LIỆU THIẾT KẾ HƯỚNG ĐỐI TƯỢNG\n(MÔ HÌNH LỚP)")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor.from_string(BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run("HỆ THỐNG QUẢN LÝ CHI TIÊU CÁ NHÂN CÓ TÍCH HỢP AI")
    r.bold = True
    r.font.size = Pt(16)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(16)
    p.add_run("Phiên bản 2.1 — As-built, schema chốt triển khai\nCập nhật theo mã nguồn ngày 02/09/2026").italic = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(35)
    p.add_run("Nhóm 21 - Thành viên nhóm 2\n").bold = True
    p.add_run("1. Đào Minh Tâm (Trưởng nhóm)\n2. Phạm Đức Nam")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(30)
    p.add_run("Thời gian thực hiện: 27/07/2026 – 27/09/2026 (9 tuần)")
    doc.add_page_break()

    # Header and footer
    header = section.header.paragraphs[0]
    header.text = "THIẾT KẾ HƯỚNG ĐỐI TƯỢNG — SỔ CHI TIÊU"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Arial"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(110, 110, 110)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Nhóm 21  |  OOD As-built v2.1  |  Trang ")
    run.font.size = Pt(8)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)

    # 1
    add_heading(doc, "1. Mô hình lớp (Class Diagram)", 1)
    add_heading(doc, "1.1. Giới thiệu mô hình lớp", 2)
    add_body(
        doc,
        "Mô hình lớp mô tả cấu trúc tĩnh đang tồn tại trong mã nguồn, gồm các lớp SQLAlchemy, "
        "thuộc tính lưu trữ và quan hệ khóa ngoại. Bản cập nhật lấy app/models, migration "
        "20260901_0001 và các router hiện hành làm chuẩn."
    )
    add_body(
        doc,
        "Hệ thống hiện có 9 lớp thực thể lưu trữ: User, Category, Transaction, Budget, "
        "SavingGoal, GoalItem, GoalTransaction, AIConversation và AIMessage. "
        "Các model là lớp dữ liệu; nghiệp vụ không được cài dưới dạng phương thức entity mà "
        "nằm chủ yếu ở Router và các hàm Core."
    )
    add_note(
        doc,
        "Điểm hiệu chỉnh",
        "Bản cũ chỉ có 8 lớp, thiếu AIConversation/AIMessage, thiếu một số cột mới và mô tả AI "
        "như định hướng. Bản này phản ánh Assistant dùng Gemini đã được triển khai. Model/bảng "
        "AIReport không có use case đã được loại bỏ trước khi chốt schema triển khai.",
        LIGHT_GREEN,
    )

    add_heading(doc, "1.2. Class Diagram tổng thể", 2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(CLASS_DIAGRAM), width=Inches(7.0))
    add_caption(doc, "Hình 1. Class Diagram as-built, đối chiếu với cấu trúc bảng hiện tại")

    add_heading(doc, "1.3. Mô tả các nhóm lớp", 2)
    add_table(
        doc,
        ["Nhóm", "Thành phần hiện tại", "Chức năng"],
        [
            ("Entity/Model", "9 lớp SQLAlchemy", "Ánh xạ dữ liệu nghiệp vụ và lịch sử hội thoại."),
            ("Schema", "user, category, transaction, budget, goal, dashboard, report, assistant", "Kiểm tra và định dạng I/O bằng Pydantic."),
            ("Router", "auth, categories, transactions, budgets, goals, dashboard, reports, assistant", "HTTP controller và phần lớn điều phối nghiệp vụ."),
            ("Core", "security, deps, finance, email, assistant_context, gemini, config", "Xác thực, tính toán dùng chung và tích hợp ngoài."),
            ("Database", "engine, SessionLocal, Base, get_db, Alembic", "Kết nối, phiên ORM và migration."),
            ("Presentation", "React + TypeScript, PWA, Windows/Tauri", "Giao diện và gọi API dùng chung."),
            ("External", "Google Gemini, Resend", "Tư vấn tài chính và gửi email đặt lại mật khẩu."),
        ],
    )

    add_heading(doc, "1.4. Quan hệ giữa các lớp", 2)
    add_table(
        doc,
        ["Quan hệ", "Bội số", "Ý nghĩa as-built"],
        [
            ("User – Category/Transaction/Budget/SavingGoal/AIConversation", "1 : 0..*", "Các bảng con có user_id; mọi truy vấn nghiệp vụ giới hạn theo user."),
            ("Category – Transaction", "1 : 0..*", "Transaction.category_id phân loại giao dịch."),
            ("Category – Budget", "1 : 0..*", "Budget.category_id áp dụng hạn mức cho danh mục."),
            ("SavingGoal – GoalItem", "1 : 0..*", "GoalItem thuộc mục tiêu; router xóa item trước khi xóa goal."),
            ("SavingGoal – GoalTransaction", "1 : 0..*", "Nhật ký nạp/rút; goal có lịch sử tiền bị chặn xóa."),
            ("AIConversation – AIMessage", "1 : 0..*", "Router xóa messages trước khi xóa conversation."),
        ],
    )

    add_heading(doc, "1.5. Bội số (Multiplicity)", 2)
    add_table(
        doc,
        ["Ký hiệu", "Ý nghĩa"],
        [
            ("1", "Chính xác một đối tượng."),
            ("0..1", "Không có hoặc có một đối tượng."),
            ("0..*", "Không có hoặc có nhiều đối tượng."),
            ("1..*", "Có ít nhất một đối tượng."),
        ],
    )
    add_body(
        doc,
        "Các model hiện khai báo ForeignKey nhưng không khai báo SQLAlchemy relationship/ondelete. "
        "Vì vậy, một số ý nghĩa vòng đời được thi hành trong Router thay vì cascade ở ORM/database."
    )

    add_heading(doc, "1.6. Vai trò của từng lớp", 2)
    add_table(
        doc,
        ["Lớp", "Vai trò"],
        [
            ("User", "Tài khoản, hồ sơ, avatar và khôi phục mật khẩu."),
            ("Category", "Danh mục income/expense với màu và icon."),
            ("Transaction", "Khoản thu/chi cụ thể; nguồn dữ liệu tổng hợp."),
            ("Budget", "Hạn mức theo danh mục và tháng/năm."),
            ("SavingGoal", "Mục tiêu, trạng thái và kết quả hoàn thành."),
            ("GoalItem", "Hạng mục con của mục tiêu."),
            ("GoalTransaction", "Nhật ký deposit/withdraw và nguồn tính current_amount."),
            ("AIConversation", "Luồng hội thoại AI của người dùng."),
            ("AIMessage", "Câu hỏi/câu trả lời cùng kỳ và evidence snapshot."),
        ],
    )

    add_heading(doc, "1.7. Nguyên tắc thiết kế hướng đối tượng", 2)
    add_bullets(
        doc,
        [
            "Entity mỏng: model tập trung vào ánh xạ dữ liệu, không gắn nhầm các endpoint thành phương thức entity.",
            "Phân tách theo module chức năng: mỗi Router/Schema/Model phục vụ một miền nghiệp vụ.",
            "Dependency Injection: Session và current_user được cấp qua Depends.",
            "Tái sử dụng có chọn lọc: finance, security, email, assistant_context và gemini được tách thành Core.",
            "Ranh giới riêng tư: Assistant chỉ nhận dữ liệu tổng hợp đã loại định danh và dữ liệu thô.",
        ],
    )

    # 2
    doc.add_page_break()
    add_heading(doc, "2. Đặc tả các lớp (Class Specification)", 1)
    add_body(
        doc,
        "Tên thuộc tính dưới đây dùng snake_case để khớp trực tiếp với mã Python và schema database. "
        "Mục “Thao tác liên quan” chỉ rõ nơi cài đặt nghiệp vụ; đây không phải phương thức của lớp ORM."
    )
    for spec in CLASS_SPECS:
        add_heading(doc, spec["title"], 2)
        add_body(doc, spec["summary"])
        p = doc.add_paragraph()
        p.add_run("Thuộc tính").bold = True
        add_table(doc, ["Thuộc tính", "Kiểu dữ liệu", "Mô tả"], spec["attrs"])
        p = doc.add_paragraph()
        p.add_run("Thao tác liên quan (Router/Core)").bold = True
        add_table(doc, ["Thao tác", "Nơi cài đặt", "Trách nhiệm"], spec["ops"])
        add_body(doc, f"Quan hệ: {spec['relations']}", bold_prefix="Quan hệ:")
        p = doc.add_paragraph()
        p.add_run("Quy tắc nghiệp vụ").bold = True
        add_bullets(doc, spec["rules"])

    add_heading(doc, "2.10. Đặc tả module Dashboard", 2)
    add_body(
        doc,
        "Dashboard không phải entity. DashboardRouter tổng hợp dữ liệu trực tiếp từ Transaction, "
        "Category, Budget, SavingGoal và GoalTransaction, đồng thời gọi compute_available_balance."
    )
    add_table(
        doc,
        ["Đầu ra", "Nguồn/tính toán"],
        [
            ("summary", "Thu, chi, net và available_balance theo kỳ."),
            ("spending_by_category", "SUM expense theo Category."),
            ("trend", "Thu/chi của 3 tháng gần nhất."),
            ("budgets", "Hạn mức, spent, remaining, usage %, status."),
            ("goals", "Tối đa 3 goal, current_amount và progress %."),
            ("recent_transactions", "Tối đa 5 giao dịch gần nhất trong kỳ."),
        ],
    )

    add_heading(doc, "2.11. Đặc tả module Reports", 2)
    add_body(
        doc,
        "ReportsRouter tạo báo cáo so sánh kỳ được chọn với tháng liền trước. Kết quả được tính xác định "
        "từ Transaction/Category và trả trực tiếp; không gọi Gemini và không lưu một bảng báo cáo riêng."
    )
    add_table(
        doc,
        ["Kết quả", "Nội dung"],
        [
            ("Summary", "Income, expense, net của hai kỳ; difference và change percentage."),
            ("Category comparison", "Amount, share, transaction_count và chênh lệch theo danh mục."),
            ("has_data", "Cho phép giao diện chọn biểu đồ hoặc empty state."),
        ],
    )

    add_heading(doc, "2.12. Đặc tả module Assistant", 2)
    add_body(
        doc,
        "AssistantRouter cung cấp context tổng hợp, quản lý lịch sử hội thoại và gửi câu hỏi tới Gemini. "
        "AssistantContext tạo evidence cho kỳ hiện tại/tháng trước; GeminiService chỉ sinh tư vấn read-only."
    )
    add_table(
        doc,
        ["API", "Trách nhiệm"],
        [
            ("GET /assistant/context", "Evidence, trạng thái có dữ liệu, ghi chú riêng tư và model."),
            ("GET /assistant/conversations", "Danh sách hội thoại theo thời gian cập nhật."),
            ("GET /assistant/conversations/{id}", "Chi tiết hội thoại và messages."),
            ("DELETE /assistant/conversations/{id}", "Xóa messages rồi xóa conversation."),
            ("POST /assistant/messages", "Gọi Gemini; lưu cặp message và evidence nếu thành công."),
        ],
    )

    # 3
    doc.add_page_break()
    add_heading(doc, "3. Thiết kế Package", 1)
    add_heading(doc, "3.1. Mục đích", 2)
    add_body(
        doc,
        "Package được mô tả theo đúng cấu trúc thư mục hiện hành, giúp truy vết từ giao diện tới API, "
        "nghiệp vụ dùng chung và dữ liệu."
    )
    add_heading(doc, "3.2. Các Package của hệ thống", 2)
    add_table(
        doc,
        ["STT", "Package", "Thành phần chính", "Chức năng"],
        [
            ("1", "frontend/src", "pages, components, auth, lib", "React UI, quản lý phiên và gọi API."),
            ("2", "app/routers", "8 router chức năng", "HTTP controller và điều phối nghiệp vụ."),
            ("3", "app/schemas", "8 nhóm Pydantic schema", "Validate/serialize request-response."),
            ("4", "app/models", "9 lớp SQLAlchemy", "Ánh xạ bảng."),
            ("5", "app/core", "security, deps, finance, email, assistant_context, gemini, config", "Logic dùng chung và tích hợp."),
            ("6", "app/database.py", "Engine, SessionLocal, Base, get_db", "Truy cập database."),
            ("7", "alembic", "migration 20260901_0001", "Quản lý schema production."),
            ("8", "uploads", "avatars", "Lưu tệp ảnh đại diện."),
        ],
    )
    add_heading(doc, "3.3. Mô tả các Package", 2)
    add_body(
        doc,
        "Presentation hỗ trợ web responsive/PWA và Windows/Tauri, cùng kết nối tới FastAPI. "
        "Router sử dụng Schema, Model và Core; Core có thể truy vấn Model qua Session hoặc gọi dịch vụ ngoài. "
        "Không có package Repository tổng quát trong mã nguồn hiện tại."
    )
    add_heading(doc, "3.4. Quan hệ giữa các Package", 2)
    add_body(
        doc,
        "Luồng chính: frontend → Router → Schema/Core/Model → Database. Riêng quên mật khẩu gọi Resend; "
        "AssistantContext đọc dữ liệu tổng hợp và GeminiService gọi Google Gemini."
    )
    add_heading(doc, "3.5. Lợi ích và giới hạn", 2)
    add_bullets(
        doc,
        [
            "Cấu trúc theo chức năng dễ tìm và phù hợp quy mô hiện tại.",
            "Schema tách validation khỏi ORM; Core giảm lặp cho nghiệp vụ dùng chung.",
            "Router vẫn chứa nhiều điều phối; khi nghiệp vụ tăng có thể tách service/repository có kiểm soát.",
        ],
    )

    # 4
    doc.add_page_break()
    add_heading(doc, "4. Thiết kế Service Layer", 1)
    add_heading(doc, "4.1. Giới thiệu", 2)
    add_body(
        doc,
        "Hệ thống dùng kiến trúc router-centric có Core services, không phải mô hình hoàn toàn không có Service Layer. "
        "CRUD và điều phối theo request nằm trong Router; logic tái sử dụng hoặc tích hợp nằm trong app/core."
    )
    add_heading(doc, "4.2. Kiến trúc xử lý nghiệp vụ", 2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ARCH_DIAGRAM), width=Inches(7.0))
    add_caption(doc, "Hình 2. Kiến trúc as-built của hệ thống")

    add_heading(doc, "4.3. Các nhóm nghiệp vụ", 2)
    add_table(
        doc,
        ["Nhóm", "Router/Core", "Chức năng hiện hành"],
        [
            ("Auth/Profile", "auth + security + email", "Tài khoản, JWT, hồ sơ, avatar, reset password."),
            ("Category", "categories", "CRUD danh mục."),
            ("Transaction", "transactions", "CRUD và lọc giao dịch."),
            ("Budget", "budgets", "CRUD, spent và is_over."),
            ("SavingGoal", "goals + finance", "Goal/item, nạp/rút, hoàn thành."),
            ("Dashboard", "dashboard + finance", "Tổng hợp màn hình tổng quan."),
            ("Reports", "reports", "So sánh tháng hiện tại/tháng trước; không AI."),
            ("Assistant", "assistant + assistant_context + gemini", "Evidence, Gemini, lịch sử hội thoại."),
        ],
    )
    add_heading(doc, "4.4. Core services dùng chung", 2)
    add_table(
        doc,
        ["Module", "Trách nhiệm"],
        [
            ("security.py", "Băm/xác minh mật khẩu; tạo/giải mã JWT."),
            ("deps.py", "Lấy current_user từ bearer token."),
            ("finance.py", "Tính available_balance từ thu, chi và dòng tiền mục tiêu."),
            ("email.py", "Gửi email đặt lại mật khẩu qua Resend."),
            ("assistant_context.py", "Tạo evidence tổng hợp, forecast và nhãn danh mục an toàn."),
            ("gemini.py", "Tạo prompt, gọi Interactions API, retry và phân loại lỗi."),
            ("config.py", "Cấu hình môi trường, CORS, trusted hosts và API docs."),
        ],
    )
    add_heading(doc, "4.5. Luồng xử lý — ví dụ thêm giao dịch", 2)
    add_numbered(
        doc,
        [
            "TransactionsPage gửi category_id, amount, txn_date và note.",
            "TransactionCreate kiểm tra amount > 0 và ngày không ở tương lai.",
            "TransactionsRouter xác thực Category thuộc current_user.",
            "Router suy ra Transaction.type từ Category.type.",
            "SQLAlchemy Session INSERT và COMMIT Transaction.",
            "API trả TransactionOut; Dashboard/Reports sẽ tổng hợp từ dữ liệu mới ở lần đọc tiếp theo.",
        ],
    )
    add_heading(doc, "4.6. Ưu điểm và điểm cần lưu ý", 2)
    add_bullets(
        doc,
        [
            "Luồng request ngắn và dễ theo dõi cho CRUD.",
            "Core services đã tách các phần phức tạp/được tái sử dụng.",
            "Một số quy tắc toàn vẹn mới ở mức Router, chưa có UniqueConstraint/cascade tương ứng trong migration.",
            "Rate limit reset password lưu trong bộ nhớ, phù hợp đồ án nhưng không chia sẻ giữa nhiều process.",
        ],
    )
    add_heading(doc, "4.7. Kết luận", 2)
    add_body(
        doc,
        "Kiến trúc hiện tại phù hợp với quy mô dự án. Nếu mở rộng, nên ưu tiên tách service theo use case "
        "có nhiều transaction boundary, thay vì thêm tầng trung gian cho mọi CRUD."
    )

    # 5
    add_heading(doc, "5. Thiết kế Database Layer", 1)
    add_heading(doc, "5.1. Mục đích", 2)
    add_body(
        doc,
        "Database Layer dùng SQLAlchemy ORM để quản lý kết nối, Session và ánh xạ 9 bảng. "
        "get_db cấp một Session cho mỗi request và luôn đóng Session trong khối finally."
    )
    add_heading(doc, "5.2. Thành phần", 2)
    add_table(
        doc,
        ["Thành phần", "Chức năng as-built"],
        [
            ("engine", "Tạo từ DATABASE_URL, bật pool_pre_ping."),
            ("SessionLocal", "autocommit=False, autoflush=False."),
            ("Base", "declarative_base cho các model."),
            ("get_db()", "Cấp và đóng Session qua Depends."),
            ("Alembic", "Migration production 20260901_0001."),
            ("schema_migrations", "Bổ sung cột tương thích khi AUTO_CREATE_SCHEMA ở development."),
        ],
    )
    add_heading(doc, "5.3. Cấu hình", 2)
    add_bullets(
        doc,
        [
            "Mặc định phát triển: sqlite:///./expense.db và check_same_thread=False.",
            "PostgreSQL được hỗ trợ qua DATABASE_URL; tiền tố postgresql:// được đổi sang postgresql+psycopg://.",
            "AUTO_CREATE_SCHEMA mặc định bật ở development và tắt ở production.",
            "API có /api/health và /api/ready; readiness thực hiện SELECT 1.",
        ],
    )
    add_heading(doc, "5.4. Quy trình ghi dữ liệu", 2)
    add_numbered(
        doc,
        [
            "Schema kiểm tra kiểu và ràng buộc đầu vào.",
            "Router xác thực quyền sở hữu và quy tắc nghiệp vụ.",
            "Tạo/cập nhật/xóa model trong Session.",
            "COMMIT; khi cần thì refresh đối tượng.",
            "Khi tích hợp ngoài lỗi trong Assistant, ROLLBACK trước khi trả lỗi.",
            "get_db đóng Session khi request kết thúc.",
        ],
    )
    add_heading(doc, "5.5. Ràng buộc và tính toàn vẹn", 2)
    add_bullets(
        doc,
        [
            "Username/email có unique index trong migration.",
            "Các quan hệ dùng ForeignKey; chưa khai báo ondelete/cascade ở schema.",
            "Một số ràng buộc tổ hợp như Category(name,type,user) và Budget(category,month,year,user) được kiểm tra ở Router, chưa có UniqueConstraint.",
            "Numeric(14,2) được dùng cho tiền, tránh Float.",
            "Avatar được lưu ở file storage; DB chỉ lưu avatar_url.",
        ],
    )

    # 6
    doc.add_page_break()
    add_heading(doc, "6. Thiết kế AI Module", 1)
    add_heading(doc, "6.1. Trạng thái triển khai", 2)
    add_body(
        doc,
        "AI Module không còn là định hướng. AssistantPage, AssistantRouter, AssistantContext, GeminiService, "
        "AIConversation và AIMessage đã có trong mã nguồn. Model mặc định là gemini-3.7-flash và có thể đổi bằng GEMINI_MODEL."
    )
    add_heading(doc, "6.2. Dữ liệu đầu vào an toàn", 2)
    add_body(
        doc,
        "AssistantContext tổng hợp kỳ hiện tại và tháng liền trước: thu, chi, net, số giao dịch, số dư khả dụng, "
        "tối đa 7 nhóm chi tiêu theo nhãn icon an toàn, ngân sách, mục tiêu và forecast."
    )
    add_bullets(
        doc,
        [
            "Không gửi username, email, internal id, ghi chú giao dịch, giao dịch thô hoặc tên mục tiêu.",
            "Không gửi tên danh mục tùy biến; ánh xạ icon sang nhãn an toàn.",
            "Câu hỏi và tối đa 10 thông điệp lịch sử dạng văn bản được dùng để tạo prompt.",
        ],
    )
    add_heading(doc, "6.3. Chức năng AI hiện hành", 2)
    add_table(
        doc,
        ["Chức năng", "Trạng thái"],
        [
            ("Xem phạm vi dữ liệu/evidence", "Đã triển khai"),
            ("Hỏi đáp tư vấn tài chính", "Đã triển khai qua Gemini"),
            ("Dự báo chi tiêu cuối tháng", "Đã triển khai trong evidence; confidence low/medium"),
            ("Lịch sử hội thoại", "Đã triển khai"),
            ("Liên kết về Reports/Transactions", "Được giao diện sử dụng để kiểm chứng dữ liệu"),
            ("AI tự sửa giao dịch/ngân sách", "Không hỗ trợ; read-only"),
            ("Báo cáo so sánh tháng", "Đã triển khai nhưng là deterministic, không dùng Gemini"),
        ],
    )
    add_heading(doc, "6.4. Luồng xử lý AI", 2)
    add_numbered(
        doc,
        [
            "Người dùng chọn kỳ và nhập câu hỏi.",
            "AssistantRouter kiểm tra hoặc tạo AIConversation.",
            "AssistantContext truy vấn và ẩn danh dữ liệu tổng hợp.",
            "GeminiService tạo system instruction, context JSON và lịch sử.",
            "Gọi Google Gemini Interactions API với retry có giới hạn.",
            "Nếu thành công, lưu cặp AIMessage và evidence_json; nếu lỗi, rollback.",
            "Trả câu trả lời, bằng chứng, độ tin cậy và nguồn điều hướng cho giao diện.",
        ],
    )
    add_heading(doc, "6.5. Kiểm soát lỗi và an toàn", 2)
    add_table(
        doc,
        ["Tình huống", "Phản hồi API"],
        [
            ("Chưa cấu hình API key", "503"),
            ("Rate limit", "429"),
            ("Timeout", "504"),
            ("Gemini không sẵn sàng", "503"),
            ("Lỗi dịch vụ/response không hợp lệ", "502"),
        ],
    )
    add_heading(doc, "6.6. Chiến lược lưu trữ kết quả AI", 2)
    add_note(
        doc,
        "As-built",
        "Câu hỏi, câu trả lời, kỳ phân tích và evidence snapshot được lưu trực tiếp trong "
        "AIConversation/AIMessage. Hệ thống không duy trì bảng báo cáo AI riêng.",
        LIGHT_GREEN,
    )

    # 7
    doc.add_page_break()
    add_heading(doc, "7. Áp dụng các nguyên tắc thiết kế hướng đối tượng", 1)
    add_heading(doc, "7.1. Đánh giá SOLID theo mức áp dụng thực tế", 2)
    add_table(
        doc,
        ["Nguyên tắc", "Mức áp dụng", "Bằng chứng/giới hạn"],
        [
            ("SRP", "Một phần", "Module chia theo miền; một số Router vẫn vừa kiểm tra, truy vấn và điều phối."),
            ("OCP", "Một phần", "Thêm router/schema/core theo module; chưa có abstraction/plugin cho mọi nghiệp vụ."),
            ("LSP", "Không phải trọng tâm", "Không có hệ phân cấp kế thừa nghiệp vụ đáng kể ngoài BaseModel/Base ORM."),
            ("ISP", "Phù hợp theo API", "Schema request/response nhỏ theo từng use case."),
            ("DIP/DI", "Có", "FastAPI Depends cấp Session và current_user; tích hợp ngoài vẫn gọi qua module cụ thể."),
        ],
    )
    add_heading(doc, "7.2. Kiến trúc tổng thể", 2)
    add_body(
        doc,
        "Kiến trúc thực tế là Presentation → Router → Schema/Core/Model → Database, kèm External Services. "
        "Bản Windows/Tauri và Web/PWA dùng chung backend FastAPI."
    )
    add_heading(doc, "7.3. Dependency Injection", 2)
    add_body(
        doc,
        "get_db và get_current_user được tiêm qua Depends vào endpoint. Cách này tách việc tạo Session/xác thực "
        "khỏi logic use case và hỗ trợ kiểm thử bằng dependency override."
    )
    add_heading(doc, "7.4. Định hướng mở rộng phù hợp", 2)
    add_bullets(
        doc,
        [
            "Bổ sung DB UniqueConstraint cho các ràng buộc tổ hợp quan trọng.",
            "Làm rõ cascade/ondelete hoặc tiếp tục điều phối vòng đời nhất quán trong service.",
            "Tách service cho Goals/Assistant nếu transaction boundary tiếp tục phức tạp.",
            "Thay rate limiter in-memory bằng shared store nếu chạy nhiều process.",
            "Nếu sau này có báo cáo AI định kỳ, thiết kế migration mới từ use case thực tế thay vì giữ bảng dự phòng.",
        ],
    )

    # 8
    doc.add_page_break()
    add_heading(doc, "8. Kết luận", 1)
    add_body(
        doc,
        "Tài liệu đã được đồng bộ với dự án as-built ngày 02/09/2026. Hệ thống có 9 lớp lưu trữ, "
        "8 router nghiệp vụ và kiến trúc router-centric kết hợp Core services. Dashboard và Reports "
        "là module tổng hợp không có entity riêng; Reports hoạt động xác định, còn Assistant dùng Gemini "
        "trên dữ liệu tổng hợp đã ẩn danh."
    )
    add_body(
        doc,
        "Thiết kế hiện tại phù hợp quy mô đồ án và đã hỗ trợ web/PWA/Windows dùng chung API. "
        "Các điểm cần theo dõi khi mở rộng là ràng buộc toàn vẹn ở database, vòng đời quan hệ, "
        "rate limiting phân tán và chiến lược migration sau khi schema production được chốt."
    )

    # Core document metadata
    doc.core_properties.title = "Thiết kế hướng đối tượng — Hệ thống Quản lý Chi tiêu Cá nhân có Tích hợp AI"
    doc.core_properties.subject = "OOD/Class Diagram as-built v2.1"
    doc.core_properties.author = "Nhóm 21"
    doc.core_properties.comments = "Cập nhật tự động theo mã nguồn hiện tại ngày 02/09/2026."

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    copy2(OUTPUT, DOWNLOAD_OUTPUT)
    return OUTPUT, DOWNLOAD_OUTPUT


if __name__ == "__main__":
    output, download_output = build_document()
    print(output)
    print(download_output)
