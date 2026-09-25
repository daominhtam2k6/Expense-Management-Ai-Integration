from pathlib import Path

from docx import Document
from docx.shared import Cm


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "docs" / "outputs"
DOCX = next(OUTPUT_DIR.glob("04_*class-specification_with_methods.docx"))
CLASS_DIAGRAM = ROOT / "docs" / "diagrams" / "uml" / "class-as-built" / "class-as-built-word.png"


def replace_text(paragraph, text):
    paragraph.clear()
    paragraph.add_run(text)


document = Document(DOCX)

replacements = {
    "Đối với Hệ thống Quản lý Chi tiêu Cá nhân có Tích hợp AI, mô hình lớp tập trung vào 9 lớp thực thể (Entity) mô tả toàn bộ dữ liệu nghiệp vụ: User, Category, Transaction, Budget, SavingGoal, GoalItem, GoalTransaction, AIConversation và AIMessage. Các lớp thực thể này được ánh xạ trực tiếp sang cơ sở dữ liệu qua ORM, và được thao tác thông qua các module xử lý nghiệp vụ (routers) tương ứng ở tầng backend.":
        "Mô hình as-built phân biệt 9 lớp thực thể ORM (User, Category, Transaction, Budget, SavingGoal, GoalItem, GoalTransaction, AIConversation và AIMessage) với các module xử lý ở tầng backend. Entity mô tả trạng thái lưu trữ; các hàm HTTP/nghiệp vụ nằm tại Router/Core và không được coi là phương thức đã cài đặt của entity.",
    "Hệ thống được thiết kế theo hướng gọn nhẹ: nghiệp vụ được xử lý ngay tại tầng tiếp nhận request (Router), không tách thêm các lớp Service/Repository trung gian. Lựa chọn này giúp hệ thống dễ triển khai và bảo trì ở quy mô hiện tại, đồng thời vẫn tuân thủ nguyên tắc Dependency Injection thông qua cơ chế Depends của FastAPI.":
        "Implementation hiện tại xử lý phần lớn điều phối tại các module Router và dùng Depends của FastAPI để nhận phiên cơ sở dữ liệu cùng người dùng hiện tại. Sơ đồ không tự tạo lớp Service/Repository chưa tồn tại và không thay đổi schema.",
    "Phần này đặc tả 9 lớp thực thể trong mô hình lớp. Các lớp ORM hiện chủ yếu chứa thuộc tính; phương thức nghiệp vụ được cài đặt tại Router/Core nhưng được trình bày dưới lớp liên quan để thể hiện đầy đủ trách nhiệm, dữ liệu vào/ra và điều kiện xử lý.":
        "Phần này đặc tả 9 entity ORM và liệt kê các thao tác Router/Core liên quan để truy vết hành vi. Các thao tác được đặt dưới mục riêng, không phải phương thức thành viên của entity. Riêng register, login, me, update_profile, forgot_password, reset_password, upload_avatar và delete_avatar nằm trong app/routers/auth.py; User chỉ lưu dữ liệu tài khoản đã đăng ký và vẫn tồn tại sau khi đăng xuất.",
    "Quản lý tài khoản, hồ sơ, xác thực, avatar và khôi phục mật khẩu.":
        "Lưu tài khoản đã đăng ký, hồ sơ, mật khẩu băm, avatar và token khôi phục. Bản ghi User tồn tại độc lập với phiên đăng nhập; truy cập giao diện không tự tạo User.",
    "b) Các phương thức": "b) Các thao tác liên quan tại Router/Core (không phải phương thức entity)",
}

for paragraph in document.paragraphs:
    text = paragraph.text.strip()
    if text in replacements:
        replace_text(paragraph, replacements[text])
    elif text.startswith("Phương thức "):
        replace_text(paragraph, text.replace("Phương thức ", "Thao tác Router/Core ", 1))

# Remove the stale manual page break before section 2.8. It produced an empty
# page after Word repaginated the updated operation tables.
for index, paragraph in enumerate(document.paragraphs[:-1]):
    if 'w:type="page"' not in paragraph._p.xml:
        continue
    following = next(
        (item.text.strip() for item in document.paragraphs[index + 1:] if item.text.strip()),
        "",
    )
    if following.startswith("2.8. Đặc tả Class AIConversation"):
        paragraph.clear()

# Replace the existing class-diagram image while preserving its paragraph position.
image_paragraph = document.paragraphs[12]
image_paragraph.clear()
image_paragraph.alignment = 1
image_paragraph.add_run().add_picture(str(CLASS_DIAGRAM), width=Cm(16.0))

# Clarify every operation table and qualify auth functions with their real module.
auth_operations = {
    "register", "login", "get_profile", "update_profile", "forgot_password",
    "reset_password", "upload_avatar", "delete_avatar",
}
for table in document.tables:
    if len(table.columns) != 2 or not table.rows:
        continue
    if table.cell(0, 0).text.strip() != "Tên":
        continue
    operation = table.cell(0, 1).text.strip()
    table.cell(0, 0).text = "Thao tác (Router/Core)"
    if operation in auth_operations:
        implementation_name = "me" if operation == "get_profile" else operation
        table.cell(0, 1).text = f"app.routers.auth.{implementation_name}"

# The role table must not imply that User performs authentication itself.
for table in document.tables:
    for row in table.rows:
        if row.cells and row.cells[0].text.strip() == "User" and len(row.cells) >= 2:
            row.cells[1].text = "Entity lưu tài khoản đã đăng ký và hồ sơ; tồn tại độc lập với phiên. Xác thực do module auth xử lý."

document.save(DOCX)
print(DOCX)
