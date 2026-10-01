from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
CAPTURE_SOURCE = ROOT / ".tmp" / "live_uc005"
FLOW_DIR = ROOT / "docs" / "diagrams" / "screenflow" / "uc005-account-management"
CAPTURE_DIR = FLOW_DIR / "captures"
FLOW_PNG = FLOW_DIR / "uc005-account-management-flow.png"
ERD_DIR = ROOT / "docs" / "diagrams" / "erd" / "screenflow-database"
ERD_PNG = ERD_DIR / "screenflow-database-erd.png"
SOURCE_DOCX = ROOT / "docs" / "outputs" / "06_GenAI_SoftwareDevelopment_screenflow_db_srs_ord_aligned.docx"
OUTPUT_DOCX = ROOT / "docs" / "outputs" / "06_GenAI_SoftwareDevelopment_screenflow_db_uc005_refined.docx"


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    filename = "arialbd.ttf" if bold else "arial.ttf"
    return ImageFont.truetype(str(Path(r"C:\Windows\Fonts") / filename), size)


def fit_image(image: Image.Image, box: tuple[int, int, int, int]) -> tuple[Image.Image, tuple[int, int]]:
    x1, y1, x2, y2 = box
    copy = image.copy()
    copy.thumbnail((x2 - x1, y2 - y1), Image.Resampling.LANCZOS)
    return copy, (x1 + (x2 - x1 - copy.width) // 2, y1 + (y2 - y1 - copy.height) // 2)


def build_screenflow() -> None:
    FLOW_DIR.mkdir(parents=True, exist_ok=True)
    CAPTURE_DIR.mkdir(parents=True, exist_ok=True)
    source_names = [
        "01-dashboard.png",
        "02-profile-panel.png",
        "04-profile-save-success.png",
        "05-password-error.png",
        "06-password-success.png",
    ]
    sources = {name: Image.open(CAPTURE_SOURCE / name).convert("RGB") for name in source_names}

    crops = {
        "01-open-account.png": sources["01-dashboard.png"].crop((0, 790, 330, 1100)),
        "02-profile-panel.png": sources["02-profile-panel.png"].crop((1015, 0, 1440, 1100)),
        "03-profile-success.png": sources["04-profile-save-success.png"].crop((1015, 0, 1440, 760)),
        "04-password-error.png": sources["05-password-error.png"].crop((1015, 620, 1440, 1100)),
        "05-password-success.png": sources["06-password-success.png"].crop((1015, 620, 1440, 1100)),
    }
    for name, crop in crops.items():
        crop.save(CAPTURE_DIR / name, optimize=True)

    canvas = Image.new("RGB", (3600, 2350), "white")
    draw = ImageDraw.Draw(canvas)
    title = font(64, True)
    subtitle = font(28)
    step_font = font(32, True)
    body = font(26)
    draw.text((1800, 70), "LUỒNG QUẢN LÝ TÀI KHOẢN — UC005", fill="#17324D", font=title, anchor="ma")
    draw.text(
        (1800, 145),
        "Ảnh chụp trực tiếp từ website production · dữ liệu minh họa biệt lập",
        fill="#607080",
        font=subtitle,
        anchor="ma",
    )

    cards = {
        1: (100, 250, 1080, 1240),
        2: (1310, 250, 2290, 1240),
        3: (2520, 250, 3500, 1240),
        4: (650, 1510, 1630, 2240),
        5: (1970, 1510, 2950, 2240),
    }
    content = {
        1: ("01-open-account.png", "Mở quản lý tài khoản", "Chọn khu vực tài khoản ở AppShell."),
        2: ("02-profile-panel.png", "Xem thông tin cá nhân", "SidePanel hiển thị avatar, hồ sơ và đổi mật khẩu."),
        3: ("03-profile-success.png", "Cập nhật hồ sơ/ảnh", "Dữ liệu hợp lệ được lưu và hiển thị thông báo thành công."),
        4: ("04-password-error.png", "Sai mật khẩu hiện tại", "Hệ thống báo lỗi, giữ panel để người sử dụng sửa lại."),
        5: ("05-password-success.png", "Đổi mật khẩu thành công", "Mật khẩu mới được chấp nhận sau khi xác thực hợp lệ."),
    }

    for number, (x1, y1, x2, y2) in cards.items():
        accent = "#D64545" if number == 4 else "#178A72" if number in (3, 5) else "#2F75B5"
        draw.rounded_rectangle((x1, y1, x2, y2), 28, fill="#F8FBFE", outline=accent, width=6)
        draw.ellipse((x1 + 24, y1 + 22, x1 + 92, y1 + 90), fill=accent)
        draw.text((x1 + 58, y1 + 56), str(number), fill="white", font=step_font, anchor="mm")
        draw.text((x1 + 115, y1 + 34), content[number][1], fill="#17324D", font=step_font)
        screenshot = Image.open(CAPTURE_DIR / content[number][0]).convert("RGB")
        fitted, position = fit_image(screenshot, (x1 + 36, y1 + 120, x2 - 36, y2 - 115))
        canvas.paste(fitted, position)
        draw.text((x1 + 38, y2 - 78), content[number][2], fill="#475569", font=body)

    def arrow(points: list[tuple[int, int]], color: str) -> None:
        draw.line(points, fill=color, width=10, joint="curve")
        ex, ey = points[-1]
        px, py = points[-2]
        if abs(ex - px) >= abs(ey - py):
            sign = 1 if ex > px else -1
            draw.polygon(((ex, ey), (ex - sign * 34, ey - 22), (ex - sign * 34, ey + 22)), fill=color)
        else:
            sign = 1 if ey > py else -1
            draw.polygon(((ex, ey), (ex - 22, ey - sign * 34), (ex + 22, ey - sign * 34)), fill=color)

    arrow([(1080, 745), (1310, 745)], "#2F75B5")
    arrow([(2290, 745), (2520, 745)], "#178A72")
    arrow([(3010, 1240), (3010, 1370), (1140, 1370), (1140, 1510)], "#D64545")
    arrow([(3010, 1240), (3010, 1510), (2460, 1510)], "#178A72")
    draw.text((1770, 1338), "Nhập mật khẩu hiện tại và mật khẩu mới", fill="#475569", font=body, anchor="mm")
    draw.text((1060, 1430), "Không hợp lệ", fill="#B42318", font=body, anchor="mm")
    draw.text((2550, 1430), "Hợp lệ", fill="#087A5B", font=body, anchor="mm")
    canvas.save(FLOW_PNG, dpi=(300, 300))


def build_clean_erd() -> None:
    ERD_DIR.mkdir(parents=True, exist_ok=True)
    canvas = Image.new("RGB", (2400, 2600), "white")
    draw = ImageDraw.Draw(canvas)
    title = font(64, True)
    table_name = font(46, True)
    field_font = font(42)
    relation_font = font(36, True)
    draw.text((1200, 65), "ERD SỔ CHI TIÊU — THIẾT KẾ MỤC TIÊU", fill="#17324D", font=title, anchor="ma")

    boxes = {
        "USERS": (780, 170, 1620, 520),
        "CATEGORIES": (40, 790, 740, 1260),
        "TRANSACTIONS": (850, 790, 1550, 1260),
        "BUDGETS": (1660, 790, 2360, 1260),
        "SAVING_GOALS": (180, 1510, 980, 1900),
        "AI_CONVERSATIONS": (1420, 1510, 2220, 1900),
        "GOAL_ITEMS": (20, 2180, 720, 2530),
        "GOAL_TRANSACTIONS": (790, 2180, 1490, 2530),
        "AI_MESSAGES": (1640, 2180, 2380, 2530),
    }
    fields = {
        "USERS": ["PK id", "UQ username_normalized", "UQ email_normalized"],
        "CATEGORIES": ["PK id", "FK user_id → USERS", "AK (id, user_id)", "name_normalized · type", "UQ owner + type + name"],
        "TRANSACTIONS": ["PK id", "FK user_id → USERS", "FK (category_id, user_id)", "amount · type · txn_date"],
        "BUDGETS": ["PK id", "FK user_id → USERS", "FK (category_id, user_id)", "month · year · limit_amount", "UQ owner + category + period"],
        "SAVING_GOALS": ["PK id", "FK user_id → USERS", "target_amount · status"],
        "GOAL_ITEMS": ["PK id", "FK goal_id → SAVING_GOALS", "name · cost · is_purchased"],
        "GOAL_TRANSACTIONS": ["PK id", "FK goal_id → SAVING_GOALS", "amount · type · txn_date · note"],
        "AI_CONVERSATIONS": ["PK id", "FK user_id → USERS", "title · created_at · updated_at"],
        "AI_MESSAGES": ["PK id", "FK conversation_id", "role · content"],
    }

    for name, (x1, y1, x2, y2) in boxes.items():
        draw.rounded_rectangle((x1, y1, x2, y2), 24, fill="#F8FBFE", outline="#2F75B5", width=5)
        draw.rounded_rectangle((x1, y1, x2, y1 + 92), 22, fill="#2F75B5")
        draw.rectangle((x1, y1 + 60, x2, y1 + 92), fill="#2F75B5")
        draw.text(((x1 + x2) // 2, y1 + 47), name, fill="white", font=table_name, anchor="mm")
        y = y1 + 125
        for value in fields[name]:
            draw.text((x1 + 26, y), value, fill="#1F2937", font=field_font)
            y += 62

    def top(name: str) -> tuple[int, int]:
        x1, y1, x2, _ = boxes[name]
        return ((x1 + x2) // 2, y1)

    def bottom(name: str) -> tuple[int, int]:
        x1, _, x2, y2 = boxes[name]
        return ((x1 + x2) // 2, y2)

    color = "#64748B"

    # USERS owns the three finance roots through one clean relationship bus.
    user = bottom("USERS")
    child_names = ("CATEGORIES", "TRANSACTIONS", "BUDGETS")
    child_centers = [top(name) for name in child_names]
    bus_y = 650
    draw.line((user, (user[0], bus_y)), fill=color, width=6)
    draw.line(((child_centers[0][0], bus_y), (child_centers[-1][0], bus_y)), fill=color, width=6)
    draw.text((user[0] + 18, user[1] + 12), "1", fill="#17324D", font=relation_font)
    for end in child_centers:
        draw.line(((end[0], bus_y), end), fill=color, width=6)
        draw.text((end[0] + 14, end[1] - 42), "0..*", fill="#17324D", font=relation_font)

    # Category ownership references are RESTRICT and share one lower bus.
    category_bottom = bottom("CATEGORIES")
    transaction_bottom = bottom("TRANSACTIONS")
    budget_bottom = bottom("BUDGETS")
    restrict_y = 1360
    draw.line((category_bottom, (category_bottom[0], restrict_y)), fill=color, width=6)
    draw.line(((category_bottom[0], restrict_y), (budget_bottom[0], restrict_y)), fill=color, width=6)
    for end in (transaction_bottom, budget_bottom):
        draw.line(((end[0], restrict_y), end), fill=color, width=6)
        draw.text((end[0] + 14, end[1] + 16), "0..*", fill="#17324D", font=relation_font)
    draw.text((category_bottom[0] + 14, category_bottom[1] + 16), "1", fill="#17324D", font=relation_font)
    draw.text((930, restrict_y - 34), "RESTRICT", fill=color, font=relation_font, anchor="mm")

    # USERS also owns goals and AI conversations; route these around the finance row.
    for name, side_x in (("SAVING_GOALS", 15), ("AI_CONVERSATIONS", 2385)):
        end = top(name)
        draw.line((user, (user[0], 590), (side_x, 590), (side_x, 1430), (end[0], 1430), end), fill=color, width=6)
        draw.text((end[0] + 14, end[1] - 42), "0..*", fill="#17324D", font=relation_font)

    # Goal children.
    goal_bottom = bottom("SAVING_GOALS")
    goal_children = (top("GOAL_ITEMS"), top("GOAL_TRANSACTIONS"))
    goal_bus_y = 2050
    draw.line((goal_bottom, (goal_bottom[0], goal_bus_y)), fill=color, width=6)
    draw.line(((goal_children[0][0], goal_bus_y), (goal_children[1][0], goal_bus_y)), fill=color, width=6)
    draw.text((goal_bottom[0] + 14, goal_bottom[1] + 16), "1", fill="#17324D", font=relation_font)
    for end in goal_children:
        draw.line(((end[0], goal_bus_y), end), fill=color, width=6)
        draw.text((end[0] + 14, end[1] - 42), "0..*", fill="#17324D", font=relation_font)

    # Conversation messages.
    conversation_bottom = bottom("AI_CONVERSATIONS")
    message_top = top("AI_MESSAGES")
    draw.line((conversation_bottom, (conversation_bottom[0], 2050), (message_top[0], 2050), message_top), fill=color, width=6)
    draw.text((conversation_bottom[0] + 14, conversation_bottom[1] + 16), "1", fill="#17324D", font=relation_font)
    draw.text((message_top[0] + 14, message_top[1] - 42), "0..*", fill="#17324D", font=relation_font)
    draw.text(
        (1200, 2575),
        "PK: khóa chính · FK: khóa ngoại · AK: khóa ứng viên · UQ: duy nhất",
        fill="#64748B",
        font=field_font,
        anchor="mm",
    )
    canvas.save(ERD_PNG, dpi=(300, 300))


def replace_picture(document: Document, part_name: str, image_path: Path, width: Cm) -> None:
    for paragraph in document.paragraphs:
        for blip in paragraph._p.xpath(".//a:blip"):
            relationship_id = blip.get(qn("r:embed"))
            if str(document.part.related_parts[relationship_id].partname) == part_name:
                for child in list(paragraph._p):
                    if child.tag != qn("w:pPr"):
                        paragraph._p.remove(child)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run().add_picture(str(image_path), width=width)
                return
    raise KeyError(part_name)


def set_text(paragraph, value: str) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(value)
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)


def update_docx() -> None:
    document = Document(SOURCE_DOCX)
    replace_picture(document, "/word/media/image22.png", FLOW_PNG, Cm(16))
    replace_picture(document, "/word/media/image26.png", ERD_PNG, Cm(16))
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text == "Hình 21. Luồng hồ sơ và ảnh đại diện":
            set_text(paragraph, "Hình 21. Luồng quản lý tài khoản UC005 trên website production")
        elif text.startswith("Từ khu vực avatar ở AppShell"):
            set_text(
                paragraph,
                "Từ khu vực tài khoản ở AppShell, SidePanel Thông tin cá nhân mở mà không làm mất màn hình nghiệp vụ. "
                "Luồng as-built cho phép cập nhật tên hiển thị, username, email, ảnh đại diện và đổi mật khẩu sau khi xác minh "
                "mật khẩu hiện tại. Website production đang hiển thị PNG/JPEG/WebP, tối đa 2 MB; đây là hành vi quan sát ngày "
                "27/09/2026, không thay thế quyết định Requirements/Security Gate. Thành công và lỗi đều hiển thị ngay trong panel.",
            )
        elif text.startswith("Sơ đồ phân biệt màu xanh"):
            set_text(
                paragraph,
                "ERD dưới đây trình bày gọn schema mục tiêu theo SRS/ORD: USERS là gốc sở hữu dữ liệu; composite FK bảo vệ "
                "ownership của TRANSACTIONS/BUDGETS với CATEGORIES; CASCADE phục vụ xóa dữ liệu active theo RQ-011, còn "
                "RESTRICT ngăn xóa danh mục đang được tham chiếu. Trạng thái triển khai chi tiết được mô tả trong các bảng sau sơ đồ.",
            )
        elif text == "Hình 24. ERD as-built và thiết kế mục tiêu":
            set_text(paragraph, "Hình 24. ERD thiết kế mục tiêu theo SRS/ORD")
    document.save(OUTPUT_DOCX)


def main() -> None:
    required = [SOURCE_DOCX] + [CAPTURE_SOURCE / name for name in (
        "01-dashboard.png", "02-profile-panel.png", "04-profile-save-success.png",
        "05-password-error.png", "06-password-success.png",
    )]
    for path in required:
        if not path.is_file():
            raise FileNotFoundError(path)
    build_screenflow()
    build_clean_erd()
    update_docx()
    print(FLOW_PNG)
    print(ERD_PNG)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
