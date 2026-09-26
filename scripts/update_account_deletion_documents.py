from copy import deepcopy
from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table


ROOT = Path(__file__).resolve().parents[1]
INPUT_DIR = Path(r"C:\Users\LENOVO\Downloads\Tai lieu du an")
OUTPUT_DIR = ROOT / "docs" / "outputs"

SRS_SOURCE = INPUT_DIR / "03_GenAI_SoftwareDevelopment_requirements-specification_actor-logout-updated.docx"
OOD_SOURCE = INPUT_DIR / "04_GenAI_SoftwareDevelopment_object-oriented-design_class-specification_final-reviewed.docx"
SRS_OUTPUT = OUTPUT_DIR / "03_GenAI_SoftwareDevelopment_requirements-specification_account-deletion-updated.docx"
OOD_OUTPUT = OUTPUT_DIR / "04_GenAI_SoftwareDevelopment_object-oriented-design_class-specification_account-deletion-design-updated.docx"

ACCOUNT_USE_CASE = ROOT / "docs/diagrams/uml/srs-use-cases/use-case-decomposition-account.png"
UC005_ACTIVITY = ROOT / "docs/diagrams/uml/srs-use-cases/uc005-profile-activity.png"
UC005_SEQUENCE = ROOT / "docs/diagrams/uml/srs-use-cases/uc005-profile-sequence.png"
CLASS_DESIGN = ROOT / "docs/diagrams/uml/class-design/class-design.png"
DELETE_ACCOUNT_SEQUENCE = ROOT / "docs/diagrams/uml/class-method-sequences/delete-account.png"


def set_update_fields(document: Document) -> None:
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def set_value(table: Table, label: str, value: str) -> None:
    for row in table.rows:
        if row.cells[0].text.strip().casefold() == label.casefold():
            row.cells[1].text = value
            return
    raise KeyError(label)


def normalize_font(document: Document) -> None:
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(13)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(13)


def replace_media(docx_path: Path, replacements: dict[str, Path]) -> None:
    with NamedTemporaryFile(delete=False, suffix=".docx") as temp:
        temp_path = Path(temp.name)
    with ZipFile(docx_path, "r") as source, ZipFile(temp_path, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            replacement = replacements.get(item.filename)
            target.writestr(item, replacement.read_bytes() if replacement else source.read(item.filename))
    temp_path.replace(docx_path)


def replace_picture_by_part(document: Document, part_name: str, image_path: Path, width: Cm) -> None:
    for paragraph in document.paragraphs:
        for blip in paragraph._p.xpath(".//a:blip"):
            relationship_id = blip.get(qn("r:embed"))
            related_part = document.part.related_parts[relationship_id]
            if str(related_part.partname) == part_name:
                for child in list(paragraph._p):
                    if child.tag != qn("w:pPr"):
                        paragraph._p.remove(child)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run().add_picture(str(image_path), width=width)
                return
    raise KeyError(part_name)


def update_srs() -> None:
    document = Document(SRS_SOURCE)
    summary = document.tables[3]
    for row in summary.rows[1:]:
        if row.cells[0].text.strip() == "UC005":
            row.cells[2].text = (
                "Xem và cập nhật hồ sơ; thay đổi ảnh đại diện; đổi mật khẩu; "
                "xóa tài khoản cùng dữ liệu cá nhân sau khi xác thực lại"
            )
            break
    else:
        raise KeyError("UC005 summary row")

    profile = next(
        table for table in document.tables
        if len(table.rows) > 1 and "UC005" in table.cell(1, 1).text
    )
    set_value(
        profile,
        "Mục đích",
        "Quản lý thông tin cá nhân, thông tin bảo mật và vòng đời tài khoản của người sử dụng.",
    )
    set_value(
        profile,
        "Mô tả",
        "Cho phép người sử dụng đã đăng nhập xem/cập nhật hồ sơ, thay đổi ảnh đại diện, đổi mật khẩu "
        "hoặc xóa tài khoản và dữ liệu cá nhân. Xóa tài khoản là chức năng con UC005e của UC005, "
        "không phải use case tổng quan độc lập. Đây là yêu cầu mục tiêu; trạng thái triển khai được kiểm chứng riêng.",
    )
    set_value(
        profile,
        "Điều kiện sau",
        "Với thao tác cập nhật hợp lệ, thông tin tương ứng được lưu và tài khoản vẫn tồn tại. "
        "Với thao tác xóa hợp lệ, dữ liệu active thuộc tài khoản bị hard delete, dữ liệu tài khoản khác giữ nguyên, "
        "client xóa phiên/cache; dữ liệu còn trong backup hết retention trong tối đa 30 ngày.",
    )
    set_value(
        profile,
        "Luồng sự kiện chính (Basic Flows)",
        "B1. Người sử dụng mở khu vực Quản lý hồ sơ.\n"
        "B2. Người sử dụng chọn xem/cập nhật hồ sơ, thay đổi ảnh đại diện, đổi mật khẩu hoặc xóa tài khoản.\n"
        "B3. Với hồ sơ/ảnh đại diện, hệ thống kiểm tra dữ liệu rồi cập nhật.\n"
        "B4. Với đổi mật khẩu, hệ thống xác minh mật khẩu hiện tại, kiểm tra và lưu mật khẩu mới dưới dạng băm.\n"
        "B5. Với xóa tài khoản, người sử dụng nhập mật khẩu hiện tại và xác nhận chính xác ‘XÓA TÀI KHOẢN’.\n"
        "B6. Hệ thống xác thực yêu cầu, hard delete toàn bộ dữ liệu active thuộc tài khoản trong một transaction rồi commit.\n"
        "B7. Hệ thống dọn avatar theo cơ chế idempotent; client xóa token, cache, bản nháp và chuyển về màn hình công khai.",
    )
    set_value(
        profile,
        "Luồng sự kiện phụ (Alternative Flows)",
        "A1. Username/email trùng → yêu cầu nhập lại.\n"
        "A2. Ảnh đại diện sai định dạng, rỗng hoặc vượt giới hạn → từ chối cập nhật.\n"
        "A3. Mật khẩu hiện tại không đúng khi đổi mật khẩu → không thay đổi dữ liệu và hiển thị lỗi chung.\n"
        "A4. Mật khẩu mới không đạt chính sách hoặc xác nhận không khớp → yêu cầu nhập lại.\n"
        "A5. Mật khẩu hoặc chuỗi xác nhận xóa không hợp lệ → giữ nguyên tài khoản, dữ liệu và phiên.\n"
        "A6. Xóa một phần dữ liệu quan hệ thất bại → rollback toàn bộ transaction; không báo xóa thành công.\n"
        "A7. Dọn avatar sau commit thất bại → tài khoản vẫn đã xóa; ghi nhận để retry idempotent.",
    )

    replace_picture_by_part(document, "/word/media/image47.png", ACCOUNT_USE_CASE, Cm(16))
    replace_picture_by_part(document, "/word/media/image14.png", UC005_ACTIVITY, Cm(16))
    replace_picture_by_part(document, "/word/media/image16.png", UC005_SEQUENCE, Cm(16))
    set_update_fields(document)
    normalize_font(document)
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == "Mô tả Use Case UC004" and any(
            "UC005" in table.cell(1, 1).text for table in document.tables if len(table.rows) > 1
        ):
            paragraph.text = "Mô tả Use Case UC005"
            paragraph.style = "Heading 3"
            break
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    document.save(SRS_OUTPUT)


def update_ood() -> None:
    document = Document(OOD_SOURCE)
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == "Quản lý tài khoản, hồ sơ, xác thực, avatar và khôi phục mật khẩu.":
            paragraph.text = (
                "Quản lý tài khoản, hồ sơ, xác thực, avatar, khôi phục mật khẩu và thiết kế mục tiêu xóa tài khoản."
            )
            break

    role_table = document.tables[4]
    for row in role_table.rows:
        if row.cells[0].text.strip() == "AccountController":
            row.cells[1].text = (
                "Điều phối đăng ký, đăng nhập, hồ sơ, avatar, khôi phục mật khẩu và xóa tài khoản theo thiết kế mục tiêu."
            )
            break

    category_heading = next(
        paragraph for paragraph in document.paragraphs
        if paragraph.text.strip() == "2.2. Đặc tả Class Category"
    )
    heading = category_heading.insert_paragraph_before("Thao tác 9: delete_account (thiết kế mục tiêu — chưa triển khai)")
    heading.style = "Heading 3"

    template_table = document.tables[13]
    table_xml = deepcopy(template_table._tbl)
    category_heading._p.addprevious(table_xml)
    operation_table = Table(table_xml, document._body)
    values = {
        "Thao tác Backend": "delete_account",
        "Mô tả": "Xóa tài khoản hiện tại và toàn bộ dữ liệu active thuộc tài khoản sau khi xác thực lại.",
        "Tham số đầu vào": "current_password; confirmation = ‘XÓA TÀI KHOẢN’.",
        "Kết quả đầu ra": "204 No Content khi transaction hoàn tất.",
        "Luồng xử lý": (
            "B1. Xác thực token, mật khẩu hiện tại và chuỗi xác nhận.\n"
            "B2. Khóa User hiện tại và hard delete dữ liệu active theo thứ tự phụ thuộc trong một transaction.\n"
            "B3. Xóa User và commit; lỗi quan hệ rollback toàn bộ.\n"
            "B4. Sau commit, dọn avatar idempotent; client xóa token/cache/bản nháp."
        ),
        "Điều kiện bắt đầu": "Đã đăng nhập; User tồn tại; mật khẩu và chuỗi xác nhận hợp lệ.",
        "Điều kiện kết thúc": (
            "Tài khoản và dữ liệu active không còn; dữ liệu người khác giữ nguyên; "
            "backup hết retention trong tối đa 30 ngày."
        ),
    }
    for label, value in values.items():
        set_value(operation_table, label, value)

    caption = category_heading.insert_paragraph_before("Biểu đồ tuần tự – delete_account (thiết kế mục tiêu)")
    caption.style = "Normal"
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph = category_heading.insert_paragraph_before()
    image_paragraph.style = "Normal"
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.add_run().add_picture(str(DELETE_ACCOUNT_SEQUENCE), width=Cm(16))
    category_heading.insert_paragraph_before()

    replace_picture_by_part(document, "/word/media/image2.png", CLASS_DESIGN, Cm(16))
    set_update_fields(document)
    normalize_font(document)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    document.save(OOD_OUTPUT)


def main() -> None:
    for required in (
        SRS_SOURCE,
        OOD_SOURCE,
        ACCOUNT_USE_CASE,
        UC005_ACTIVITY,
        UC005_SEQUENCE,
        CLASS_DESIGN,
        DELETE_ACCOUNT_SEQUENCE,
    ):
        if not required.is_file():
            raise FileNotFoundError(required)
    update_srs()
    update_ood()
    print(SRS_OUTPUT)
    print(OOD_OUTPUT)


if __name__ == "__main__":
    main()
