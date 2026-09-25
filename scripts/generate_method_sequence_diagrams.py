import re
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r"C:\Users\LENOVO\Downloads\Tai lieu du an\04_GenAI_SoftwareDevelopment_object-oriented-design_class-specification.docx")
OUT_DIR = ROOT / "docs/diagrams/uml/class-method-sequences"

GROUPS = {
    "auth": {"register", "login", "get_profile", "update_profile", "forgot_password", "reset_password", "upload_avatar", "delete_avatar"},
    "categories": {"list_categories", "create_category", "update_category", "delete_category"},
    "transactions": {"list_transactions", "create_transaction", "update_transaction", "delete_transaction"},
    "budgets": {"list_budgets", "create_budget", "update_budget", "delete_budget", "compute_spent"},
    "goals": {"list_goals", "create_goal", "get_goal_detail", "update_goal", "delete_goal", "complete_goal", "add_item", "update_item", "delete_item", "list_goal_transactions", "deposit", "withdraw", "compute_current"},
    "assistant": {"list_conversations", "get_conversation", "create_conversation", "delete_conversation", "ask_assistant", "message_out"},
}

UI_NAMES = {
    "register": "Màn hình Đăng ký",
    "login": "Màn hình Đăng nhập",
    "forgot_password": "Màn hình Quên mật khẩu",
    "reset_password": "Màn hình Đặt lại mật khẩu",
    "get_profile": "Màn hình Hồ sơ",
    "update_profile": "Màn hình Hồ sơ",
    "upload_avatar": "Màn hình Hồ sơ",
    "delete_avatar": "Màn hình Hồ sơ",
}


def clean(text: str) -> str:
    return " ".join(text.replace("\n", " ").split())


def slugify(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def wrap(text: str, width: int = 28) -> str:
    text = clean(text).replace(":", " -").replace("→", "–")
    words, lines, current = text.split(), [], []
    for word in words:
        if current and len(" ".join(current + [word])) > width:
            lines.append(" ".join(current))
            current = [word]
        else:
            current.append(word)
    if current:
        lines.append(" ".join(current))
    return "\\n".join(lines)


def split_steps(flow: str) -> list[str]:
    numbered = [clean(x) for x in re.split(r"\s*\d+\)\s*", flow) if clean(x)]
    result = []
    for step in numbered:
        parts = re.split(r";\s+|\.\s+(?=[A-ZĐẢÁÀÂĂÊÔƠƯ])", step)
        for part in parts:
            part = clean(part).rstrip(".")
            if part:
                result.append(part)
    return result


def group_for(operation: str) -> str:
    return next(group for group, names in GROUPS.items() if operation in names)


def target_for(step: str, group: str) -> str:
    lower = step.lower()
    if "gemini" in lower:
        return "Gemini"
    if "email" in lower and any(word in lower for word in ("gửi", "resend")):
        return "Resend"
    if any(word in lower for word in ("jwt", "bcrypt", "băm", "hash token", "giải mã jwt")):
        return "Core"
    if any(word in lower for word in ("tìm ", "lọc ", "lấy ", "đếm ", "lưu ", "commit", "flush", "xóa ", "cập nhật", "tạo user", "tạo category", "tạo transaction", "tạo deposit", "tạo withdraw")):
        return "DB"
    if any(word in lower for word in ("available_balance", "current_amount", "tính", "sum ", "đầu/cuối tháng", "evidence", "chuẩn hóa dữ liệu")):
        return "Core"
    return "Router"


document = Document(SOURCE)
OUT_DIR.mkdir(parents=True, exist_ok=True)
manifest = []

for table in document.tables:
    if len(table.columns) != 2 or len(table.rows) < 7 or table.cell(0, 0).text.strip() != "Tên":
        continue
    values = {row.cells[0].text.strip(): row.cells[1].text.strip() for row in table.rows}
    operation = clean(values["Tên"])
    group = group_for(operation)
    interface = UI_NAMES[operation] if operation in UI_NAMES else {
        "categories": "Màn hình Danh mục",
        "transactions": "Màn hình Giao dịch",
        "budgets": "Màn hình Ngân sách",
        "goals": "Màn hình Mục tiêu tiết kiệm",
        "assistant": "Màn hình Trợ lý AI",
    }[group]
    steps = split_steps(values.get("Luồng xử lý", ""))
    slug = slugify(operation)
    lines = [
        "@startuml",
        f"title BIỂU ĐỒ TUẦN TỰ — {operation}",
        "skinparam backgroundColor white",
        "skinparam defaultFontName Arial",
        "skinparam defaultFontSize 24",
        "skinparam dpi 300",
        "skinparam shadowing false",
        "skinparam sequence {",
        "  ArrowColor #345B52",
        "  LifeLineBorderColor #299D78",
        "  ParticipantBorderColor #299D78",
        "  ParticipantBackgroundColor #F4FBF8",
        "}",
        'actor "Người sử dụng" as User',
        f'boundary "{interface}" as UI',
        'control "Backend\\n(Router + Core)" as Backend',
        'database "CSDL" as DB',
    ]
    lines.extend([
        f"User -> UI: Chọn {operation}",
        f"UI -> Backend: Gửi yêu cầu {operation}",
        "activate Backend",
    ])
    for step in steps:
        target = target_for(step, group)
        if target == "DB":
            lines.append(f"Backend -> DB: {wrap(step)}")
            lines.append("DB --> Backend: Kết quả")
        else:
            lines.append(f"Backend -> Backend: {wrap(step)}")
    outcome = clean(values.get("Điều kiện kết thúc", ""))
    lines.extend([
        "alt Xử lý thành công",
        f"  Backend --> UI: {wrap(outcome or 'Trả kết quả thành công')}",
        "  UI --> User: Hiển thị kết quả",
        "else Dữ liệu không hợp lệ hoặc lỗi xử lý",
        "  Backend --> UI: Mã lỗi và thông báo phù hợp",
        "  UI --> User: Hiển thị lỗi, giữ dữ liệu cần thiết",
        "end",
        "deactivate Backend",
        "@enduml",
        "",
    ])
    (OUT_DIR / f"{slug}.puml").write_text("\n".join(lines), encoding="utf-8")
    manifest.append((operation, slug, interface, len(steps)))

(OUT_DIR / "manifest.tsv").write_text(
    "operation\tslug\tinterface\tsteps\n" +
    "\n".join(f"{op}\t{slug}\t{interface}\t{count}" for op, slug, interface, count in manifest) + "\n",
    encoding="utf-8",
)
(OUT_DIR / "render-notes.md").write_text(
    "# Sequence Diagram cho thao tác Router/Core\n\n"
    f"- Nguồn: `{SOURCE}` và mã nguồn dự án được phản ánh trong đặc tả lớp.\n"
    "- Trạng thái: as-built ở mức thành phần; thao tác được nhóm theo entity để truy vết nhưng thuộc Router/Core.\n"
    "- Có 40 Sequence Diagram, font Arial 24 để đọc rõ khi chèn rộng tối đa 16 cm trong Word, nền trắng, render PlantUML local.\n"
    "- Mỗi hình thống nhất 4 lifeline: Người sử dụng, giao diện cụ thể, Backend (Router + Core) và CSDL. Tích hợp ngoài được mô tả trong thông điệp Backend để giữ hình gọn.\n",
    encoding="utf-8",
)
print(f"Generated {len(manifest)} sequence diagrams")
