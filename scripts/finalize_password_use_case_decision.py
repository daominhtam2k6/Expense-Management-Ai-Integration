from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "docs/outputs/03_GenAI_SoftwareDevelopment_requirements-specification.docx"
ACCOUNT_DIAGRAM = ROOT / "docs/diagrams/uml/srs-use-cases/use-case-decomposition-account.png"


def set_value(table, label, value):
    for row in table.rows:
        if row.cells[0].text.strip().casefold() == label.casefold():
            row.cells[1].text = value
            return
    raise KeyError(label)


document = Document(DOCX)

summary = document.tables[3]
for row in summary.rows[1:]:
    if row.cells[0].text.strip() == "UC005":
        row.cells[2].text = (
            "Sửa tên hiển thị, username, email; tải lên hoặc xóa ảnh đại diện; "
            "đổi mật khẩu khi đang đăng nhập"
        )
        break

profile = next(
    table for table in document.tables
    if len(table.rows) > 1 and "UC005" in table.cell(1, 1).text
)
set_value(profile, "Mục đích", "Quản lý thông tin cá nhân và thông tin bảo mật của tài khoản.")
set_value(
    profile,
    "Mô tả",
    "Cho phép người sử dụng đã đăng nhập thay đổi tên hiển thị, username, email, "
    "ảnh đại diện hoặc đổi mật khẩu. Đổi mật khẩu là chức năng con của Quản lý hồ sơ, "
    "không phải use case tổng quan độc lập.",
)
set_value(
    profile,
    "Điều kiện sau",
    "Thông tin hồ sơ, ảnh đại diện hoặc mật khẩu được cập nhật theo thao tác hợp lệ; "
    "tài khoản vẫn giữ nguyên.",
)
set_value(
    profile,
    "Luồng sự kiện chính (Basic Flows)",
    "B1. Người sử dụng mở khu vực Quản lý hồ sơ.\n"
    "B2. Người sử dụng chọn sửa thông tin, thay đổi ảnh đại diện hoặc đổi mật khẩu.\n"
    "B3. Với thông tin hồ sơ/ảnh đại diện, hệ thống kiểm tra định dạng và tính duy nhất rồi cập nhật.\n"
    "B4. Với đổi mật khẩu, người sử dụng nhập mật khẩu hiện tại, mật khẩu mới và xác nhận mật khẩu mới.\n"
    "B5. Hệ thống xác minh mật khẩu hiện tại, kiểm tra mật khẩu mới và lưu mật khẩu mới dưới dạng băm.\n"
    "B6. Hệ thống thông báo cập nhật thành công.",
)
set_value(
    profile,
    "Luồng sự kiện phụ (Alternative Flows)",
    "A1. Username/email trùng → yêu cầu nhập lại.\n"
    "A2. Ảnh đại diện sai định dạng, rỗng hoặc vượt giới hạn → từ chối cập nhật.\n"
    "A3. Mật khẩu hiện tại không đúng → không thay đổi mật khẩu và hiển thị lỗi chung.\n"
    "A4. Mật khẩu mới không đạt chính sách hoặc xác nhận không khớp → yêu cầu nhập lại.",
)

for paragraph in document.paragraphs:
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)
for table in document.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(13)

document.save(DOCX)

# The five decomposition figures are the last inserted media items; image47 is Account.
with NamedTemporaryFile(delete=False, suffix=".docx") as temp:
    temp_path = Path(temp.name)
with ZipFile(DOCX, "r") as source, ZipFile(temp_path, "w", ZIP_DEFLATED) as target:
    for item in source.infolist():
        data = ACCOUNT_DIAGRAM.read_bytes() if item.filename == "word/media/image47.png" else source.read(item.filename)
        target.writestr(item, data)
temp_path.replace(DOCX)
print(DOCX)
