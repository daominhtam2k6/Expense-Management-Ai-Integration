import re
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(r"C:\Users\LENOVO\Downloads\Tai lieu du an\04_GenAI_SoftwareDevelopment_object-oriented-design_class-specification.docx")
OUTPUT = ROOT / "docs/outputs/04_GenAI_SoftwareDevelopment_object-oriented-design_class-specification.docx"
DIAGRAM_DIR = ROOT / "docs/diagrams/uml/class-method-sequences"
CLASS_DIAGRAM = ROOT / "docs/diagrams/uml/class-design/class-design.png"


def clean(text: str) -> str:
    return " ".join(text.replace("\n", " ").split())


def slugify(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def split_steps(flow: str) -> list[str]:
    numbered = [clean(x) for x in re.split(r"\s*\d+\)\s*", flow) if clean(x)]
    result = []
    for step in numbered:
        for part in re.split(r";\s+|\.\s+(?=[A-ZĐẢÁÀÂĂÊÔƠƯ])", step):
            part = clean(part).rstrip(".")
            if part:
                result.append(part)
    return result


def set_font(run):
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")


document = Document(SOURCE)

# Keep the overview tables consistent with the hybrid design diagram.
group_table = document.tables[0]
group_row = group_table.add_row()
group_row.cells[0].text = "Điều khiển use case"
group_row.cells[1].text = "AccountController, FinancialController, GoalController, AnalysisController, AssistantController"

component_table = document.tables[1]
control_row = component_table.add_row()
control_row.cells[0].text = "Control (thiết kế)"
control_row.cells[1].text = "AccountController, FinancialController, GoalController, AnalysisController, AssistantController"
control_row.cells[2].text = "Tổ chức các thao tác Router/Core hiện có thành trách nhiệm điều phối ở mức thiết kế."

role_table = document.tables[4]
for row in role_table.rows[1:]:
    if row.cells[0].text.strip() == "User":
        row.cells[1].text = "Tài khoản đã đăng ký, hồ sơ, mật khẩu băm và ảnh đại diện."
for name, role in [
    ("AccountController", "Điều phối đăng ký, đăng nhập, hồ sơ, avatar và khôi phục mật khẩu."),
    ("FinancialController", "Điều phối danh mục, giao dịch và ngân sách."),
    ("GoalController", "Điều phối mục tiêu, hạng mục và lịch sử nạp/rút."),
    ("AnalysisController", "Tổng hợp Dashboard và báo cáo so sánh."),
    ("AssistantController", "Điều phối hội thoại, dữ liệu tổng hợp và câu trả lời AI."),
]:
    row = role_table.add_row()
    row.cells[0].text = name
    row.cells[1].text = role
for paragraph in document.paragraphs:
    text = paragraph.text.strip()
    if text.startswith("Đối với Hệ thống Quản lý Chi tiêu Cá nhân"):
        paragraph.clear()
        paragraph.add_run(
            "Class Diagram kết hợp hiện trạng dữ liệu với thiết kế hướng đối tượng. "
            "Chín entity giữ các thuộc tính, khóa và quan hệ đang được lưu trữ; các phương thức được tổ chức theo trách nhiệm của lớp. "
            "Những thao tác điều phối nhiều đối tượng được đặt ở lớp control tương ứng."
        )
    elif text.startswith("Hệ thống được thiết kế theo hướng gọn nhẹ"):
        paragraph.clear()
        paragraph.add_run(
            "Các lớp control trong sơ đồ là cách tổ chức trách nhiệm ở mức thiết kế, được truy vết từ Router/Core hiện tại. "
            "Tên lớp không khẳng định implementation đã tách thành các class Python giống hệt sơ đồ."
        )
    elif text == "Biểu đồ gồm 9 lớp thực thể chính:":
        paragraph.clear()
        paragraph.add_run(
            "Biểu đồ gồm 9 lớp thực thể và 5 lớp điều khiển, thể hiện thuộc tính, phương thức, khóa PK/FK, quan hệ và bội số."
        )
    elif text == "2. Đặc tả Class (Class Specification)":
        paragraph.clear()
        paragraph.add_run("2. Đặc tả lớp và thao tác liên quan")
    elif text.startswith("Phần này đặc tả 9 lớp thực thể"):
        paragraph.clear()
        paragraph.add_run(
            "Phần này đặc tả chi tiết 9 lớp thực thể và nhóm các thao tác Backend theo đối tượng nghiệp vụ chịu ảnh hưởng. "
            "Bảng thuộc tính bám theo dữ liệu đã triển khai; bảng thao tác và Sequence Diagram mô tả trách nhiệm thiết kế cùng luồng phối hợp với Backend."
        )
    elif text == "b) Các phương thức":
        paragraph.clear()
        paragraph.add_run("b) Các thao tác liên quan và biểu đồ tuần tự")
    elif text.startswith("Phương thức "):
        paragraph.clear()
        paragraph.add_run(text.replace("Phương thức ", "Thao tác ", 1))

# Replace the overview with the approved design-level class model.
class_diagram_paragraph = document.paragraphs[12]
class_diagram_paragraph.clear()
class_diagram_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
class_diagram_paragraph.add_run().add_picture(str(CLASS_DIAGRAM), width=Cm(16.0))

inserted = 0
for table in list(document.tables):
    if len(table.columns) == 4 and table.cell(0, 0).text.strip() == "Tên":
        for row in table.rows[1:]:
            name = row.cells[0].text.strip()
            if name == "id":
                row.cells[0].text = "id (PK)"
            elif name in {"user_id", "category_id", "goal_id", "conversation_id"}:
                row.cells[0].text = f"{name} (FK)"
    if len(table.columns) != 2 or len(table.rows) < 7 or table.cell(0, 0).text.strip() != "Tên":
        continue
    operation = table.cell(0, 1).text.strip()
    table.cell(0, 0).text = "Thao tác Backend"
    for row in table.rows:
        if row.cells[0].text.strip() == "Luồng xử lý":
            steps = split_steps(row.cells[1].text)
            row.cells[1].text = "\n".join(f"B{i}. {step}." for i, step in enumerate(steps, 1))
            break

    image_path = DIAGRAM_DIR / f"{slugify(operation)}.png"
    if not image_path.exists():
        raise FileNotFoundError(image_path)
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_with_next = True
    caption.add_run(f"Biểu đồ tuần tự – {operation}")
    picture = document.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    with Image.open(image_path) as image:
        w, h = image.size
    width_cm = min(16.0, 14.5 * w / h)
    picture.add_run().add_picture(str(image_path), width=Cm(width_cm))
    table._tbl.addnext(caption._p)
    caption._p.addnext(picture._p)
    inserted += 1

for style in document.styles:
    if style.type == 1:
        style.font.name = "Times New Roman"
        style.font.size = Pt(13)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
for paragraph in document.paragraphs:
    for run in paragraph.runs:
        set_font(run)
for table in document.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    set_font(run)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
document.save(OUTPUT)
print(f"{OUTPUT}\nInserted sequence diagrams: {inserted}")
